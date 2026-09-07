import sys

with open('kpi_reports.py', 'r', encoding='utf-8') as f:
    content = f.read()

new_func = """
def generate_yearly_excel(df, year):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"TongKet_{year}"
    
    # Headers
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(df.columns))
    ws.cell(row=1, column=1, value=f"BẢNG TỔNG KẾT KPI CẢ NĂM {year}")
    
    font_default = Font(name='Times New Roman', size=13)
    font_bold = Font(name='Times New Roman', size=13, bold=True)
    ws.cell(row=1, column=1).font = Font(name='Times New Roman', size=16, bold=True)
    ws.cell(row=1, column=1).alignment = Alignment(horizontal='center', vertical='center')
    
    thin = Side(border_style="thin", color="000000")
    border = Border(top=thin, left=thin, right=thin, bottom=thin)
    
    # Column headers
    for col_idx, col_name in enumerate(df.columns, 1):
        c = ws.cell(row=3, column=col_idx, value=col_name)
        c.font = font_bold
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = border
        c.fill = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")
    
    # Data rows
    for row_idx, row_data in enumerate(df.values, 4):
        for col_idx, val in enumerate(row_data, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=val)
            c.font = font_default
            c.border = border
            c.alignment = Alignment(horizontal='center', vertical='center')
    
    # Adjust column widths
    from openpyxl.utils import get_column_letter
    for col_idx in range(1, len(df.columns) + 1):
        col_letter = get_column_letter(col_idx)
        if col_idx == 1: ws.column_dimensions[col_letter].width = 25
        elif col_idx == 2: ws.column_dimensions[col_letter].width = 25
        else: ws.column_dimensions[col_letter].width = 12
        
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
"""

if "generate_yearly_excel" not in content:
    with open('kpi_reports.py', 'a', encoding='utf-8') as f:
        f.write(new_func)
    print("Patched kpi_reports.py OK")
else:
    print("Function already exists.")
