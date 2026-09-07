import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Title
title = doc.add_heading('KỊCH BẢN THUYẾT TRÌNH BÁO CÁO BAN LÃNH ĐẠO', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph('Chủ đề: Phần Mềm Quản Lý Công Việc & Đánh Giá KPI - DMT Group')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].bold = True

slides = [
    (
        "SLIDE 1: TRANG BÌA",
        "Kính thưa Ban Lãnh đạo và các Anh/Chị quản lý,\nHôm nay tôi xin phép được giới thiệu về phần mềm mới giúp công ty chúng ta số hóa công tác Quản lý công việc và Đánh giá KPI.\nMục tiêu lớn nhất của dự án này rất rõ ràng: *Mọi công việc đều phải được lưu trữ thành dữ liệu, và kết quả đánh giá phải dựa trên minh chứng thực tế.* Chúng ta sẽ chuyển đổi sang một hệ thống quản trị mang tính thực chất và khách quan hơn."
    ),
    (
        "SLIDE 2: THỰC TRẠNG HIỆN NAY",
        "Trước tiên, xin phép cùng nhìn lại cách chúng ta đang thực hiện đánh giá KPI hiện nay, có 4 khó khăn chính:\n1. Mất nhiều thời gian: Đến mỗi kỳ báo cáo, các phòng ban phải mất rất nhiều công sức để tổng hợp file Word, Excel.\n2. Thiếu sự phân hóa: Kết quả đánh giá cuối kỳ thường có xu hướng giống nhau, chưa thực sự làm nổi bật được những cá nhân làm việc hiệu quả vượt trội.\n3. Đánh giá chủ quan: Việc chấm điểm đôi khi phải dựa vào cảm nhận và trí nhớ của quản lý, rất khó để bám sát toàn bộ tiến độ công việc trong tháng.\n4. Thiếu lịch sử đối chiếu: Khi cần xem lại một công việc bị trễ tiến độ, chúng ta mất nhiều thời gian để tra cứu lý do và kết quả thực tế."
    ),
    (
        "SLIDE 3: GIẢI PHÁP PHẦN MỀM MỚI",
        "Để giải quyết những khó khăn trên, phần mềm mới mang đến 3 giá trị cốt lõi:\n- Số hóa toàn diện: Chuyển hoàn toàn quy trình giao việc và báo cáo từ file văn bản lên một hệ thống duy nhất.\n- Đánh giá có minh chứng: Điểm KPI sẽ được hệ thống tự động ghi nhận dựa trên file kết quả đính kèm thực tế, đảm bảo tính khách quan tuyệt đối.\n- Theo dõi liên tục: Ban Lãnh đạo có thể xem Bảng tổng quan để nắm bắt tiến độ công việc mọi lúc, mọi nơi, không cần đợi đến cuối kỳ."
    ),
    (
        "SLIDE 4: CẤU TRÚC HỆ THỐNG CHUYÊN BIỆT",
        "Để hệ thống vận hành trơn tru nhưng không làm người dùng bị 'ngợp', chúng tôi thiết kế cấu trúc phần mềm chia làm 2 nhóm chính: Công cụ dùng chung và Tính năng phân quyền.\n- Về công cụ dùng chung:\n  1. Chúng ta có Bảng Tổng Quan (Dashboard) giúp Ban Lãnh đạo và mọi người nhìn bao quát toàn bộ tiến độ, nhận diện ngay các rủi ro trễ hạn.\n  2. Kế đến là Sổ tay Hướng dẫn - một cẩm nang trực tuyến chung để bất kỳ ai cũng có thể tra cứu quy trình và cách dùng phần mềm.\n- Về tính năng phân quyền:\n  3. Với Nhân Viên: Giao diện được làm tối giản nhất có thể, chỉ tập trung vào việc tạo mới, nhận việc và báo cáo tiến độ.\n  4. Với Quản Lý: Hệ thống mở rộng thêm các tính năng chuyên sâu để xem xét vướng mắc khách quan, duyệt việc và chốt kết quả KPI."
    ),
    (
        "SLIDE 5: QUY TRÌNH TẠO VIỆC & TÁI TẠO ĐỊNH KỲ",
        "Về quy trình tạo công việc trên hệ thống, chúng ta sẽ có luồng thao tác rất rõ ràng như sau:\n- Đầu tháng: Các nhân sự sẽ chủ động lên hệ thống để đăng ký các 'Công việc định kỳ' của mình. Nếu trong tháng có việc phát sinh thêm thì đăng ký bổ sung.\n- Công việc giao ban: Riêng với những công việc mang tính chỉ đạo, được sếp giao đột xuất trong các cuộc họp giao ban, nhân sự được giao việc hoàn toàn có thể lên hệ thống đăng ký bổ sung ngay sau khi cuộc họp kết thúc.\n\nVà để nhân sự không thấy phiền khi phải đăng ký nhiều việc, tính năng tạo việc được thiết kế vô cùng thông minh.\nHệ thống đã thiết lập sẵn các danh mục từ điển. Ví dụ, khi mọi người chọn tên một phòng ban, hệ thống sẽ tự động lọc (sort) và hiển thị chính xác danh sách nhân sự của phòng ban đó. Chúng ta chỉ việc click chọn thay vì gõ tay, giúp chuẩn hóa dữ liệu ngay từ đầu.\n\nĐặc biệt nhất là tính năng Tái tạo định kỳ. Ví dụ: Tháng nào bộ phận HCNS cũng có việc 'Tính lương'. Sang tháng mới, nhân viên không cần tạo lại từ đầu, chỉ cần bấm 'Tái tạo', phần mềm sẽ tự động sao chép toàn bộ thông tin cũ sang, giảm thiểu tối đa thời gian nhập liệu."
    ),
    (
        "SLIDE 6: BÁO CÁO, CHỈNH SỬA TIẾN ĐỘ & XỬ LÝ VƯỚNG MẮC",
        "Quy trình báo cáo của nhân viên sẽ diễn ra như sau:\n- Khi hoàn thành: Nhập trạng thái và bắt buộc đính kèm file kết quả.\n- Tính năng Sửa đổi: Trong quá trình thao tác, nếu nhân viên lỡ cập nhật nhầm trạng thái hoặc đính kèm nhầm file, mọi người hoàn toàn có thể sử dụng chức năng Sửa đổi để cập nhật lại thông tin cho chính xác.\n- Khi gặp vướng mắc: Nhân viên cần ghi nhận rõ nguyên nhân là do bản thân (Chủ quan) hay do các yếu tố bên ngoài (Khách quan - ví dụ: chờ phản hồi từ đối tác).\n- Với nguyên nhân khách quan, Quản lý sẽ trực tiếp vào xem xét mức độ công việc đã thực hiện để duyệt phần trăm hoàn thành (ví dụ 50% hoặc 80%), đảm bảo quyền lợi chính đáng cho nhân sự."
    ),
    (
        "SLIDE 7: BẢNG THEO DÕI TIẾN ĐỘ CHUNG",
        "Tất cả trạng thái công việc từ các phòng ban sẽ được tổng hợp tự động về Bảng theo dõi tiến độ chung.\nNgay khi nhìn vào giao diện, Ban Lãnh đạo và Quản lý sẽ thấy ngay bức tranh toàn cảnh thông qua các con số thống kê rất trực quan: Tổng số việc là bao nhiêu, tiến độ Đã xong, Đang làm. Và quan trọng nhất là hệ thống tự động bóc tách số lượng các việc đang Trễ hạn hoặc Vướng mắc (được làm nổi bật bằng màu đỏ cảnh báo) để chúng ta lưu tâm.\n\nKhông chỉ dừng lại ở các con số, hệ thống còn chia sẵn thành các mục chức năng rất tiện lợi ở bên dưới:\n- Mục Công việc tới hạn (và trễ hạn): Giúp quản lý xem ngay danh sách những việc nào đang 'cháy' tiến độ để đốc thúc xử lý kịp thời.\n- Mục Báo cáo giao ban: Tự động tổng hợp các công việc trọng tâm, trở thành tài liệu sống để trình chiếu và thảo luận ngay trong các buổi họp giao ban mà không cần ai phải làm file báo cáo riêng.\n- Cùng với bộ công cụ Lọc (Sort) linh hoạt theo từng nhân sự hay dự án... chúng ta hoàn toàn chấm dứt cảnh phải kéo chuột dò tìm thủ công qua hàng ngàn dòng Excel như trước kia."
    ),
    (
        "SLIDE 8: QUYỀN QUẢN LÝ - DUYỆT & XÓA CÔNG VIỆC",
        "Trong quy trình này, cấp Quản lý đóng vai trò như một chốt chặn kiểm soát dữ liệu:\n- Duyệt việc khách quan (DVKQ): Thao tác này rất đơn giản. Đến cuối tháng, Quản lý chỉ cần truy cập vào mục 'Duyệt Việc Khách Quan' trên phần mềm. Hệ thống đã tự động gom sẵn tất cả các báo cáo vướng mắc khách quan lại. Quản lý chỉ việc xem trong danh sách đó có nhân sự thuộc phòng ban mình hay không, đánh giá khối lượng công việc thực tế của họ và duyệt mức độ hoàn thành (ví dụ 50%, 80%). Tỷ lệ % này sẽ được liên kết trực tiếp vào điểm KPI cuối kỳ, đảm bảo quyền lợi công bằng.\n- Tính năng Xóa công việc: Nếu phát hiện nhân sự đăng ký nhầm công việc, tạo trùng lặp hoặc đưa lên những việc không hợp lý, Quản lý có thể can thiệp. Tính năng Xóa công việc được phân quyền chỉ dành riêng cho Quản lý. Điều này giúp tránh tình trạng nhân viên tự ý xóa bừa bãi các công việc đang làm dở dang để trốn tránh trách nhiệm, đảm bảo tính toàn vẹn của dữ liệu toàn công ty."
    ),
    (
        "SLIDE 9: QUY TRÌNH CHẤM KPI CUỐI KỲ",
        "Nhờ dữ liệu được cập nhật hàng ngày, công tác chấm KPI cuối kỳ trở nên rất đơn giản.\nHệ thống sẽ tự động tính toán Điểm KPI đề xuất dựa trên tiến độ và minh chứng.\nCác Trưởng bộ phận và Tổ KPI không cần tổng hợp thủ công. Quản lý chỉ cần kiểm tra lại bảng điểm hệ thống xuất ra, đối chiếu các trường hợp đặc biệt và xác nhận kết quả cuối cùng."
    ),
    (
        "SLIDE 10: CƠ CẤU ĐIỂM & THƯỞNG / PHẠT",
        "Về cơ cấu điểm, hệ thống phân bổ 70% cho Công việc định kỳ và 30% cho Công việc giao ban.\nNgoài ra, chúng ta có chính sách Thưởng/Phạt minh bạch trên hệ thống:\n- Cộng điểm: Dành cho việc hoàn thành xuất sắc hoặc có sáng kiến mang lại hiệu quả thiết thực.\n- Trừ điểm: Áp dụng khi vi phạm nội quy, đi trễ, hoặc sai sót trong nghiệp vụ.\nBộ phận HCNS và Quản lý sẽ rà soát các tiêu chí này trước khi kết thúc kỳ đánh giá."
    ),
    (
        "SLIDE 11: XẾP LOẠI & TRẢ LƯƠNG",
        "Tổng điểm cuối cùng sẽ quy ra các mức xếp loại như sau:\n- Loại D (Kém - dưới 71 điểm): Không đạt yêu cầu.\n- Loại C và B: Mức lương KPI tương ứng là 60% và 80%.\n- Loại A (Giỏi - trên 91 điểm): Nhận 100% lương KPI.\n- Loại A* (Xuất sắc - trên 100 điểm): Đạt mức 110–120% lương, nhằm khích lệ tinh thần làm việc vượt trội."
    ),
    (
        "SLIDE 12: ĐÁNH GIÁ NHÂN SỰ XUẤT SẮC (A*)",
        "Tiêu chí đạt A* đòi hỏi sự nỗ lực rất lớn và được Ban Lãnh đạo xét duyệt định kỳ hàng quý.\nĐiểm nổi bật của hệ thống là khả năng cộng dồn: Nếu nhân sự duy trì được thành tích A* trong phần lớn các quý, kết quả cuối năm của họ sẽ được hệ thống ghi nhận ở mức A* năm. Điều này giúp chúng ta đánh giá đúng và giữ chân những cá nhân có phong độ cống hiến ổn định."
    ),
    (
        "SLIDE 13: CÔNG NGHỆ AI KIỂM TRA MÔ TẢ CÔNG VIỆC",
        "Một bước tiến về công nghệ của hệ thống là việc tích hợp Trí tuệ nhân tạo (AI) để hỗ trợ theo dõi Mô tả công việc (JD).\n\nChính vì vậy, việc các phòng ban xây dựng và cập nhật bản JD ban đầu tương đối chính xác là rất quan trọng. Dựa trên bản JD chuẩn do Trưởng bộ phận mô tả, AI sẽ hỗ trợ Tổ KPI giám sát xem những công việc nhân viên đang làm thực tế có bị lệch pha so với chức năng nhiệm vụ hay không.\n\nTất nhiên, chúng tôi hiểu rằng có một số phòng ban đặc thù, công việc biến đổi liên tục và khó ghi cụ thể. Nhưng các anh/chị yên tâm, AI đủ thông minh để nhận diện ngữ nghĩa xem công việc đó thực chất nằm trong nhóm chức năng nào để phân loại. Hệ thống và Tổ KPI sẽ chỉ đưa ra báo động khi nhân sự đó làm những việc lệch quá nhiều, quá xa so với chuyên môn trong JD.\n\nMục đích cuối cùng của tính năng này hoàn toàn là để theo dõi và giám sát xem các đầu việc nhân viên đưa ra có hợp lý hay không, qua đó giúp Ban Lãnh đạo và Quản lý đảm bảo nguyên tắc 'đúng người, đúng việc'."
    ),
    (
        "SLIDE 14: LỘ TRÌNH ÁP DỤNG & CHẠY THỬ NGHIỆM",
        "Để đưa phần mềm vào vận hành hiệu quả, chúng ta sẽ đi theo một lộ trình triển khai rất rõ ràng và thận trọng:\n1. Thành lập Tổ KPI chuyên trách: Chúng ta sẽ lập ra một Tổ KPI quy tụ đại diện từ các mảng cốt lõi của công ty gồm: Tài chính, Dự án, Kỹ thuật và Nhân sự. Việc có đầy đủ chuyên môn sẽ giúp Tổ KPI nắm bắt và thấu hiểu được tính chất công việc của toàn Group. Tổ này sẽ chịu trách nhiệm họp lại để rà soát và chốt điểm KPI định kỳ theo tháng, theo quý.\n2. Dữ liệu đầu vào: Các phòng ban tiến hành rà soát lại JD và khởi tạo dữ liệu lên phần mềm để chuẩn hóa.\n3. Thử nghiệm (Pilot) từ Quý IV: Chúng ta sẽ bắt đầu chạy thử nghiệm phần mềm ngay từ Quý IV tới đây. Trong giai đoạn này, Tổ dự án sẽ liên tục thu thập các ý kiến đóng góp, phản hồi từ tất cả mọi người dùng để tinh chỉnh và hoàn thiện phần mềm sao cho trơn tru nhất.\n4. Vận hành chính thức: Sau thời gian thử nghiệm, toàn bộ dữ liệu công việc và KPI sẽ được chuyển sang chạy 100% trên phần mềm. Yếu tố quyết định thành công lúc này là sự nghiêm túc cập nhật dữ liệu của tất cả chúng ta."
    ),
    (
        "SLIDE 15: LỜI KẾT",
        "Tóm lại, giải pháp phần mềm này không chỉ giúp tiết kiệm thời gian vận hành mà còn mang lại môi trường làm việc minh bạch, nơi năng lực và sự cống hiến được đo lường bằng số liệu thực tế.\nXin trân trọng cảm ơn Ban Lãnh đạo đã lắng nghe. Tiếp theo, tôi xin phép trình diễn thao tác trực tiếp trên phần mềm để làm rõ hơn các tính năng."
    )
]

for slide_title, content in slides:
    head = doc.add_heading(slide_title, level=2)
    head.runs[0].font.color.rgb = docx.shared.RGBColor(0, 102, 204)
    p = doc.add_paragraph(content)
    
doc.save('Kich_Ban_Thuyet_Trinh_DMT.docx')
print("Saved to Kich_Ban_Thuyet_Trinh_DMT.docx")
