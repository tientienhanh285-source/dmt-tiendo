import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_shading(cell, color_hex):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    """Set inner margins (padding) of a cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_callout_border(cell, color_hex="0066FF", sz="24"):
    """Set thick left border and none for others."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:top w:val="none"/>'
                        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>'
                        f'<w:bottom w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)

def set_table_borders(table, color_hex="CBD5E1"):
    """Set clean light borders for a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
                        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
                        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color_hex}"/>'
                        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>'
                        f'<w:insideV w:val="none"/>'
                        f'<w:left w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'</w:tblBorders>')
    tblPr.append(borders)

def create_handbook():
    doc = docx.Document()
    
    # Page Setup: A4, 2cm margins
    sections = doc.sections
    for section in sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.79)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("DMT GROUP — SỔ TAY HƯỚNG DẪN SỬ DỤNG TRÊN ĐIỆN THOẠI (SMARTPHONE)")
        hrun.font.name = "Segoe UI"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Ban Hành chính Nhân sự • Lưu hành nội bộ • Quý IV/2026")
        frun.font.name = "Segoe UI"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)

    # Base Colors
    NAVY = RGBColor(10, 37, 64)       # #0A2540
    ROYAL = RGBColor(0, 102, 255)     # #0066FF
    GOLD = RGBColor(217, 119, 6)      # #D97706
    GREEN = RGBColor(16, 185, 129)    # #10B981
    DARK = RGBColor(30, 41, 59)       # #1E293B
    GRAY = RGBColor(100, 116, 139)    # #64748B

    # Helper functions for text
    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = "Segoe UI"
        run.font.size = Pt(20)
        run.bold = True
        run.font.color.rgb = NAVY
        return p

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(18)
        run = p.add_run(text)
        run.font.name = "Segoe UI"
        run.font.size = Pt(12)
        run.bold = True
        run.font.color.rgb = ROYAL
        return p

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Segoe UI"
        run.font.size = Pt(14)
        run.bold = True
        run.font.color.rgb = NAVY
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Segoe UI"
        run.font.size = Pt(12)
        run.bold = True
        run.font.color.rgb = ROYAL
        return p

    def add_heading_3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Segoe UI"
        run.font.size = Pt(11)
        run.bold = True
        run.font.color.rgb = DARK
        return p

    def add_paragraph(text="", bold_prefix="", italic=False, color=DARK):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Segoe UI"
            r_pre.font.size = Pt(10.5)
            r_pre.bold = True
            r_pre.font.color.rgb = color
        if text:
            r_text = p.add_run(text)
            r_text.font.name = "Segoe UI"
            r_text.font.size = Pt(10.5)
            r_text.italic = italic
            r_text.font.color.rgb = color
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Segoe UI"
            r_pre.font.size = Pt(10.5)
            r_pre.bold = True
            r_pre.font.color.rgb = DARK
        r_text = p.add_run(text)
        r_text.font.name = "Segoe UI"
        r_text.font.size = Pt(10.5)
        r_text.font.color.rgb = DARK
        return p

    def add_callout(text, bold_prefix="LƯU Ý QUAN TRỌNG: ", color_hex="0066FF", bg_hex="F0F7FF"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.6)
        set_cell_shading(cell, bg_hex)
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        set_callout_border(cell, color_hex=color_hex, sz="24")
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix + " ")
            r_pre.font.name = "Segoe UI"
            r_pre.font.size = Pt(10)
            r_pre.bold = True
            r_pre.font.color.rgb = ROYAL if color_hex == "0066FF" else (GOLD if color_hex == "D97706" else GREEN)
        r_text = p.add_run(text)
        r_text.font.name = "Segoe UI"
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = DARK
        
        # spacer
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)

    def add_mobile_image(img_path, caption):
        if os.path.exists(img_path):
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            cell.width = Inches(3.2)
            set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
            set_cell_shading(cell, "F8FAFC")
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run()
            run.add_picture(img_path, width=Inches(3.0))
            
            cp = cell.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(4)
            r_cap = cp.add_run(f"📱 {caption}")
            r_cap.font.name = "Segoe UI"
            r_cap.font.size = Pt(9.5)
            r_cap.bold = True
            r_cap.font.color.rgb = ROYAL
            
            # spacer
            sp = doc.add_paragraph()
            sp.paragraph_format.space_after = Pt(6)

    # ------------------ BẮT ĐẦU NỘI DUNG TÀI LIỆU ------------------

    # COVER BOX
    cover_tbl = doc.add_table(rows=1, cols=1)
    cover_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_cell = cover_tbl.cell(0, 0)
    c_cell.width = Inches(6.6)
    set_cell_shading(c_cell, "0A2540")
    set_cell_margins(c_cell, top=260, bottom=260, left=240, right=240)
    
    cp1 = c_cell.paragraphs[0]
    cp1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = cp1.add_run("TẬP ĐOÀN DMT GROUP • BAN HÀNH CHÍNH NHÂN SỰ\n")
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(245, 158, 11)
    r_sub.bold = True
    
    r_main = cp1.add_run("SỔ TAY HƯỚNG DẪN THAO TÁC TRÊN ĐIỆN THOẠI DI ĐỘNG\n")
    r_main.font.name = "Segoe UI"
    r_main.font.size = Pt(16)
    r_main.font.color.rgb = RGBColor(255, 255, 255)
    r_main.bold = True
    
    r_title = cp1.add_run("HỆ THỐNG QUẢN LÝ TIẾN ĐỘ CÔNG VIỆC & ĐÁNH GIÁ KPI")
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(13)
    r_title.font.color.rgb = RGBColor(56, 189, 248)
    r_title.bold = True

    cp2 = c_cell.add_paragraph()
    cp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp2.paragraph_format.space_before = Pt(8)
    r_desc = cp2.add_run("Tài liệu hướng dẫn trực quan dành cho Toàn thể Cán bộ Nhân viên & Cấp Quản lý\nÁp dụng trong giai đoạn Chạy thử nghiệm Quý IV/2026")
    r_desc.font.name = "Segoe UI"
    r_desc.font.size = Pt(9.5)
    r_desc.font.color.rgb = RGBColor(203, 213, 225)
    r_desc.italic = True

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # PHẦN MỞ ĐẦU
    add_heading_1("LỜI NÓI ĐẦU: VÌ SAO SỬ DỤNG TRÊN ĐIỆN THOẠI DI ĐỘNG?")
    add_paragraph("Kính gửi các anh chị Trưởng Ban và toàn thể Cán bộ Nhân viên Tập đoàn DMT Group,")
    add_paragraph("Để tạo điều kiện thuận lợi tối đa cho công tác theo dõi công việc ngoài công trường, trong các cuộc họp và khi đi công tác xa, Ban HCNS đã tối ưu hóa toàn bộ giao diện Hệ thống Quản lý Tiến độ & KPI trên nền tảng di động (Smartphone). Giờ đây, chỉ với chiếc điện thoại thông minh kết nối Internet, các anh chị có thể tạo việc, báo cáo tiến độ và duyệt việc mọi lúc mọi nơi chỉ trong vài giây.")

    add_callout(
        "💡 MẸO TIỆN DỤNG CỰC HAY: Đưa phần mềm ra Màn hình chính như một App độc lập!\n"
        "• Trên iPhone (Safari): Bấm nút Chia sẻ (biểu tượng hình vuông có mũi tên chỉ lên) ➔ Chọn 'Thêm vào MH chính' (Add to Home Screen).\n"
        "• Trên Android (Chrome/Cốc Cốc): Bấm dấu 3 chấm góc trên bên phải ➔ Chọn 'Cài đặt ứng dụng' hoặc 'Thêm vào Màn hình chính'.\n"
        "Sau khi cài, bạn sẽ có ngay icon ứng dụng DMT trên màn hình điện thoại, bấm vào là mở ứng dụng ngay không cần gõ lại địa chỉ web!",
        bold_prefix="MẸO DÙNG TIỆN LỢI:", color_hex="10B981", bg_hex="ECFDF5"
    )

    # 5 NGUYÊN TẮC CỐT LÕI
    add_heading_1("5 NGUYÊN TẮC VẬN HÀNH BẮT BUỘC (CŨ vs MỚI)")
    add_paragraph("Để hệ thống vận hành trơn tru và phát huy tối đa giá trị, toàn thể CBNV và Quản lý cần nắm vững 5 nguyên tắc cốt lõi sau:")

    tbl_rules = doc.add_table(rows=6, cols=3)
    tbl_rules.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_rules)
    
    headers = ["Nguyên tắc", "Cách làm cũ (Thủ công / Zalo)", "Nguyên tắc mới trên Phần mềm"]
    col_widths = [Inches(1.5), Inches(2.5), Inches(2.6)]
    
    # Format header
    hdr_cells = tbl_rules.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_shading(hdr_cells[i], "0A2540")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.bold = True
        run.font.name = "Segoe UI"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(255, 255, 255)

    rules_data = [
        ("1. Nguồn gốc việc", "Tự nghĩ việc rời rạc, làm theo thói quen, không rõ gắn với mục tiêu nào.", "Bắt buộc gắn với Kế hoạch năm và Chỉ tiêu Ban (Chọn Dự án ➔ Chọn Chỉ tiêu ➔ Nhập tên việc)."),
        ("2. Trách nhiệm hai chiều", "Coi báo cáo là việc của nhân viên nộp cho sếp; sếp giao việc miệng qua Zalo dễ quên.", "Cả Sếp và Nhân viên đều có tài khoản và cùng làm việc; Sếp vừa phân rã chỉ tiêu vừa tự tạo & làm việc của sếp."),
        ("3. Cập nhật Real-time", "Dồn toa cuối tháng mới cuống cuồng ngồi nhớ lại; vướng mắc để dồn cuối kỳ mới vỡ lở.", "Cập nhật liên tục ngay khi làm; có nguy cơ trễ hạn hoặc vướng mắc phải báo sớm để Ban hỗ trợ tháo gỡ."),
        ("4. Minh chứng thực tế", "Khai báo 'Đã xong' bằng miệng hoặc giấy tờ chung chung; duyệt theo cảm tính.", "Bắt buộc đính kèm File / Link minh chứng sản phẩm; nghiệm thu dựa trên kết quả thực tế ngoài đời."),
        ("5. Phân biệt 2 cấp độ", "Ngộ nhận: Thấy xong việc cá nhân là tưởng Chỉ tiêu Ban đã xong, dẫn đến số liệu ảo.", "Rạch ròi 2 cấp độ: Việc cá nhân xong KHÔNG mặc nhiên là xong Chỉ tiêu; Quản lý phải trực tiếp nghiệm thu Master View.")
    ]

    for row_idx, data in enumerate(rules_data, start=1):
        row_cells = tbl_rules.rows[row_idx].cells
        bg_col = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            row_cells[col_idx].width = col_widths[col_idx]
            set_cell_shading(row_cells[col_idx], bg_col)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=100, right=100)
            p = row_cells[col_idx].paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.name = "Segoe UI"
            run.font.size = Pt(9.5)
            if col_idx == 0:
                run.bold = True
                run.font.color.rgb = NAVY
            elif col_idx == 1:
                run.font.color.rgb = RGBColor(225, 29, 72)
            else:
                run.font.color.rgb = RGBColor(16, 185, 129)
                run.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # PHẦN 1: ĐĂNG NHẬP
    add_heading_1("PHẦN 1: HƯỚNG DẪN ĐĂNG NHẬP TRÊN ĐIỆN THOẠI")
    add_paragraph("Giao diện đăng nhập trên điện thoại được tích hợp gọn gàng trong thanh trượt (Sidebar Menu):")
    add_bullet("Bước 1: Chạm vào biểu tượng mũi tên '>' (hoặc dấu 3 gạch) ở góc trên cùng bên trái màn hình để mở thanh bên.", "Mở Menu: ")
    add_bullet("Bước 2: Chọn Công ty / Đơn vị thành viên của bạn (Tập đoàn DMT, Cienco, v.v.).", "Chọn Đơn vị: ")
    add_bullet("Bước 3: Chọn Phòng ban / Ban chuyên môn trực thuộc.", "Chọn Phòng ban: ")
    add_bullet("Nếu là Nhân viên: Chạm vào ô 'Tên nhân viên' và chọn đúng tên của mình trong danh sách (Không cần mật khẩu).", "Đăng nhập Nhân viên: ")
    add_bullet("Nếu là Quản lý (Trưởng/Phó Ban): Chọn vai trò 'Quản lý / Lãnh đạo' và nhập mật khẩu quản lý: quanly123.", "Đăng nhập Quản lý: ")
    add_bullet("Bước 4: Chạm vào dấu 'X' hoặc chạm tay vào vùng màn hình bên phải để đóng menu và vào màn hình làm việc chính.", "Vào màn hình chính: ")

    add_mobile_image("screenshots_mobile/mobile_01_dang_nhap.png", "Hình 1: Giao diện đăng nhập và chọn nhân sự trên điện thoại")

    # PHẦN 2: DASHBOARD
    add_heading_1("PHẦN 2: MÀN HÌNH TỔNG QUAN (DASHBOARD CÁ NHÂN)")
    add_paragraph("Sau khi chọn đúng tên, hệ thống lập tức hiển thị bức tranh toàn cảnh công việc của bạn trong tháng:")
    add_bullet("Tổng số công việc: Toàn bộ đầu việc bạn tự đăng ký hoặc được cấp trên phân công.", "1. Thẻ màu xanh dương — ")
    add_bullet("Hoàn thành: Số công việc đã làm xong và được Quản lý duyệt nghiệm thu.", "2. Thẻ màu xanh lá — ")
    add_bullet("Đang thực hiện: Các công việc đang trong quá trình tác nghiệp.", "3. Thẻ màu vàng/cam — ")
    add_bullet("Cảnh báo quá hạn: Những việc đã vượt quá Deadline cần tập trung xử lý ngay.", "4. Thẻ màu đỏ — ")
    add_paragraph("Thanh tiến độ % ở phía dưới cho biết tỷ lệ hoàn thành kế hoạch tháng tính đến thời điểm hiện tại.")

    add_mobile_image("screenshots_mobile/mobile_02_dashboard_nhanvien.png", "Hình 2: Màn hình Tổng quan cá nhân và các thẻ chỉ số KPI")

    # PHẦN 3: DANH SÁCH CÔNG VIỆC
    add_heading_1("PHẦN 3: XEM VÀ TRA CỨU DANH SÁCH CÔNG VIỆC")
    add_paragraph("Cuộn màn hình điện thoại xuống dưới để xem chi tiết từng công việc trong danh sách:")
    add_bullet("Mỗi công việc hiển thị đầy đủ: Tên Dự Án, Mục Tiêu Ban, Tên công việc cụ thể, Hạn chót (Deadline) và Trạng thái.", "Hiển thị chi tiết: ")
    add_bullet("Trên màn hình điện thoại, bạn có thể vuốt ngang trên bảng dữ liệu để xem thêm các cột: Người phối hợp, Ghi chú, Link minh chứng...", "Thao tác vuốt ngang: ")
    add_bullet("Màu sắc trạng thái giúp nhận diện cực nhanh: Đã hoàn thành (Xanh lá), Đang làm (Xanh dương), Vướng mắc khách quan (Vàng), Quá hạn (Đỏ).", "Nhận diện màu sắc: ")

    add_mobile_image("screenshots_mobile/mobile_03_danh_sach_cong_viec.png", "Hình 3: Danh sách công việc cá nhân hiển thị trên màn hình điện thoại")

    # PHẦN 4: THANH ĐIỀU HƯỚNG
    add_heading_1("PHẦN 4: CHUYỂN TRANG BẰNG THANH ĐIỀU HƯỚNG MENU")
    add_paragraph("Bất cứ lúc nào, bạn chỉ cần chạm vào biểu tượng Menu góc trên bên trái để chuyển qua lại giữa các trang chức năng:")
    add_bullet("Trang 1: 📊 Tổng Quan — Xem các chỉ số KPI và biểu đồ cá nhân.", "Menu chính: ")
    add_bullet("Trang 2: 📋 Danh Sách Công Việc — Bảng tra cứu, lọc và sắp xếp công việc theo thời gian/tiến độ.", "")
    add_bullet("Trang 3: ✍️ Cập Nhật Tiến Độ — Đăng ký việc mới và báo cáo kết quả hoàn thành.", "")
    add_bullet("Trang 4: ⭐ Đánh Giá KPI — Bảng điểm tháng, điểm thưởng (+5đ), xếp loại A/B/C/D và thưởng A* (110%-120% lương).", "")
    add_bullet("Trang 5: 👥 Quản Lý & Phê Duyệt — Dành cho Trưởng/Phó Ban duyệt việc và nghiệm thu chỉ tiêu.", "")

    add_mobile_image("screenshots_mobile/mobile_04_sidebar_menu.png", "Hình 4: Menu điều hướng các tính năng trên điện thoại")

    # PHẦN 5: TẠO VIỆC & CẬP NHẬT TIẾN ĐỘ
    add_heading_1("PHẦN 5: ĐĂNG KÝ CÔNG VIỆC MỚI & CẬP NHẬT TIẾN ĐỘ (THAO TÁC TRỌNG TÂM)")
    add_paragraph("Đây là tính năng các anh chị em sẽ sử dụng thường xuyên nhất hàng ngày và hàng tuần:")

    add_heading_2("5.1. Trình tự 3 bước tạo công việc mới chuẩn:")
    add_paragraph("Theo quy định của Tập đoàn, mọi công việc đều phải xuất phát từ mục tiêu chung:")
    add_bullet("Bước 1: Chạm chọn 'Tên Dự Án' tương ứng trong danh mục thả xuống.", "1. Chọn Dự án: ")
    add_bullet("Bước 2: Chạm chọn 'Chỉ Tiêu Ban' mà công việc của bạn hướng tới.", "2. Chọn Chỉ tiêu: ")
    add_bullet("Bước 3: Nhập tên công việc cụ thể vào ô 'Tên công việc' (Ví dụ: 'Lập tờ trình thẩm tra thiết kế đợt 2').", "3. Nhập tên việc: ")

    add_heading_2("5.2. Tính năng đột phá: 'TÁI TẠO ĐỊNH KỲ' (1-Click) sang tháng mới:")
    add_paragraph("Đối với các công việc thường xuyên lặp lại hàng tháng (như: Báo cáo thuế, tính lương, kiểm tra thiết bị công trường định kỳ, trực ban...):")
    add_bullet("Đầu tháng mới, bạn chỉ cần bấm đúng 1 nút: 'TÁI TẠO ĐỊNH KỲ'.", "Tiết kiệm 90% thời gian: ")
    add_bullet("Toàn bộ danh mục công việc chuẩn của tháng trước sẽ tự động nhân bản sang tháng mới với đầy đủ Dự án và Chỉ tiêu, bạn không cần phải gõ lại từng việc!", "")

    add_mobile_image("screenshots_mobile/mobile_06_form_nhap_lieu.png", "Hình 5: Form nhập liệu công việc chuẩn trên điện thoại")

    add_heading_2("5.3. Báo cáo hoàn thành & Đính kèm Minh chứng:")
    add_paragraph("Khi hoàn thành công việc, hãy thực hiện các thao tác sau ngay trên điện thoại:")
    add_bullet("Chuyển trạng thái từ 'Đang làm' sang 'Hoàn thành'.", "1. Chọn trạng thái: ")
    add_bullet("Gõ tóm tắt kết quả (ví dụ: 'Đã hoàn thành bàn giao hồ sơ nghiệm thu đợt 1').", "2. Ghi chú kết quả: ")
    add_bullet("Dán đường link tài liệu minh chứng (Link Google Drive, OneDrive hoặc Link ảnh hiện trường công việc).", "3. Link minh chứng: ")

    add_heading_2("5.4. Cơ chế bảo vệ: Nút 'VƯỚNG MẮC KHÁCH QUAN':")
    add_paragraph("Nếu công việc bị chậm tiến độ vì nguyên nhân ngoài tầm kiểm soát (Ví dụ: Chờ thủ tục cơ quan nhà nước, đối tác chậm gửi hồ sơ, mưa bão công trường dừng thi công...):")
    add_bullet("Chạm chọn 'Vướng mắc khách quan' và ghi rõ lý do vào ô giải trình.", "Thao tác tự bảo vệ: ")
    add_bullet("Quản lý sẽ xem xét khối lượng thực tế bạn đã làm và duyệt tỷ lệ hoàn thành (50%, 80%...).", "Quản lý thẩm định: ")
    add_bullet("Kết quả: Bạn được bảo toàn điểm số KPI, TUYỆT ĐỐI KHÔNG BỊ TRỪ ĐIỂM OAN!", "Ý nghĩa nhân văn: ")

    add_mobile_image("screenshots_mobile/mobile_07_form_nut_luu.png", "Hình 6: Ô nhập minh chứng, nút vướng mắc và nút Lưu dữ liệu")

    add_callout(
        "⚠️ LƯU Ý KỸ THUẬT QUAN TRỌNG TRÊN ĐIỆN THOẠI:\n"
        "Sau khi điền xong thông tin, bạn hãy chạm nút 'LƯU CÔNG VIỆC' ĐÚNG 1 LẦN DUY NHẤT và đợi khoảng 1-2 giây cho hệ thống ghi nhận và hiện thông báo màu xanh 'Thành công'.\n"
        "Tuyệt đối KHÔNG chạm đúp (bấm liên tục 2 lần) vào nút Lưu để tránh hệ thống tạo ra 2 công việc trùng nhau!",
        bold_prefix="LƯU Ý TRÁNH TRÙNG VIỆC:", color_hex="D97706", bg_hex="FFFBEB"
    )

    # PHẦN 6: DÀNH CHO QUẢN LÝ
    add_heading_1("PHẦN 6: HƯỚNG DẪN DÀNH CHO CẤP QUẢN LÝ (TRƯỞNG / PHÓ BAN)")
    add_paragraph("Dành riêng cho các anh chị Lãnh đạo và Trưởng/Phó Ban khi thao tác duyệt việc trên điện thoại:")

    add_mobile_image("screenshots_mobile/mobile_11_quan_ly_dashboard.png", "Hình 7: Màn hình quản lý tổng quan tiến độ phòng ban")

    add_heading_2("6.1. Giám sát toàn diện trên điện thoại:")
    add_paragraph("Khi đăng nhập bằng quyền Quản lý (mật khẩu quanly123), màn hình Dashboard sẽ hiển thị tổng hợp tiến độ của tất cả nhân sự trong phòng ban. Bạn có thể nhanh chóng nắm bắt ai đang làm tốt, ai đang có việc chậm trễ ngay cả khi đang ngồi họp.")

    add_heading_2("6.2. Thao tác Duyệt công việc (Approve Task) — Cấp độ 1:")
    add_bullet("Nhân viên báo Hoàn thành ➔ Công việc sẽ xuất hiện trong danh sách 'Chờ duyệt'.", "Quy trình duyệt: ")
    add_bullet("Quản lý chạm vào Link minh chứng để kiểm tra sản phẩm thực tế.", "Kiểm tra minh chứng: ")
    add_bullet("Bấm nút 'Duyệt' để xác nhận hoàn thành và cộng điểm KPI cho nhân sự. Nếu chưa đạt, bấm yêu cầu bổ sung.", "Phê duyệt: ")

    add_heading_2("6.3. Nghiệm thu Chỉ tiêu Ban (Close Objective trên Master View) — Cấp độ 2:")
    add_bullet("Dù các việc nhánh của nhân viên đã đạt 100%, Chỉ tiêu lớn của Ban vẫn KHÔNG tự động đóng.", "Chốt chặn chất lượng: ")
    add_bullet("Trưởng ban là người trực tiếp thẩm định tổng thể chất lượng công việc theo Kế hoạch năm và bấm nút 'Nghiệm thu Chỉ tiêu' để hoàn tất.", "Trách nhiệm Trưởng ban: ")

    # PHẦN 7: FAQ
    add_heading_1("PHẦN 7: GIẢI ĐÁP 6 CÂU HỎI THƯỜNG GẶP KHI DÙNG TRÊN ĐIỆN THOẠI")

    faqs = [
        ("Câu 1: Lỡ bấm nhầm hoặc tạo sai tên công việc thì sửa thế nào trên điện thoại?",
         "Bạn chỉ cần vào trang 'Cập nhật tiến độ', chạm chọn 'Sửa công việc đã có', chọn đúng tên công việc cần sửa và cập nhật lại nội dung. Quản lý cũng có quyền xóa các việc bị tạo trùng hoặc sai sót."),
        ("Câu 2: Không tải trực tiếp được file nặng từ điện thoại lên thì đính kèm minh chứng bằng cách nào?",
         "Rất đơn giản! Bạn chỉ cần tải tài liệu lên Google Drive, OneDrive của công ty hoặc chụp ảnh lưu trên album, sau đó sao chép (Copy) đường link chia sẻ và dán vào ô 'Link minh chứng'. Cách này vừa nhanh, vừa không tốn dung lượng bộ nhớ điện thoại."),
        ("Câu 3: Một công việc có 2 người cùng làm thì ai là người báo cáo?",
         "Người được giao vai trò 'Người chủ trì' sẽ là người cập nhật chính và chịu trách nhiệm về kết quả. 'Người phối hợp' vẫn được ghi nhận công sức và đánh giá đóng góp tương ứng."),
        ("Câu 4: Quên mật khẩu đăng nhập của Quản lý thì liên hệ ai?",
         "Mật khẩu mặc định của Quản lý là: quanly123. Nếu cần cấp lại hoặc đổi mật khẩu, vui lòng liên hệ ngay Ban HCNS để được hỗ trợ trong 1 phút."),
        ("Câu 5: Đến cuối tháng mà việc chưa xong do lý do khách quan thì giải trình thế nào?",
         "Bạn chỉ cần chọn trạng thái 'Vướng mắc khách quan' và ghi rõ lý do. Quản lý sẽ thẩm định và duyệt bảo lưu điểm số, tuyệt đối bạn không bị trừ điểm oan."),
        ("Câu 6: Trong Quý IV này nếu tôi thao tác sai hoặc công việc bị chậm thì có bị trừ lương không?",
         "HOÀN TOÀN KHÔNG! Ban Lãnh đạo và Ban HCNS cam kết Quý IV/2026 là giai đoạn CHẠY THỬ (PILOT) để toàn thể anh chị em làm quen hệ thống. Tuyệt đối không áp dụng chế tài trừ điểm hay trừ lương trong giai đoạn này!")
    ]

    for q, a in faqs:
        add_callout(a, bold_prefix=f"❓ {q}\n👉 TRẢ LỜI: ", color_hex="0066FF", bg_hex="F8FAFC")

    # FOOTER BOX
    foot_tbl = doc.add_table(rows=1, cols=1)
    foot_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    ft_cell = foot_tbl.cell(0, 0)
    ft_cell.width = Inches(6.6)
    set_cell_shading(ft_cell, "0A2540")
    set_cell_margins(ft_cell, top=180, bottom=180, left=200, right=200)
    
    fp = ft_cell.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c = fp.add_run("BAN HÀNH CHÍNH NHÂN SỰ TẬP ĐOÀN DMT GROUP\n")
    r_c.font.name = "Segoe UI"
    r_c.font.size = Pt(11)
    r_c.font.color.rgb = RGBColor(245, 158, 11)
    r_c.bold = True
    
    r_s = fp.add_run("🤝 Luôn đồng hành 24/7 cùng các anh chị em — Hoàn thiện hệ thống vì người dùng!\nHotline / Zalo hỗ trợ kỹ thuật: Ban HCNS DMT")
    r_s.font.name = "Segoe UI"
    r_s.font.size = Pt(9.5)
    r_s.font.color.rgb = RGBColor(255, 255, 255)

    output_path = "Huong_Dan_Su_Dung_Mobile_DMT.docx"
    doc.save(output_path)
    print(f"Successfully generated Word file: {output_path} ({os.path.getsize(output_path):,} bytes)")

if __name__ == "__main__":
    create_handbook()
