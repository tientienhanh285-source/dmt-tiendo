import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_slide(prs, title_text, content_bullets, note_text="", include_image_placeholder=False, image_caption=""):
    # Chọn layout có tiêu đề và nội dung (thường là layout số 1)
    slide_layout = prs.slide_layouts[1] 
    slide = prs.slides.add_slide(slide_layout)
    
    # Thiết lập tiêu đề
    title = slide.shapes.title
    title.text = title_text
    for paragraph in title.text_frame.paragraphs:
        paragraph.font.name = 'Arial'
        paragraph.font.bold = True
        paragraph.font.color.rgb = RGBColor(0, 51, 102) # Màu xanh dương đậm
    
    # Thiết lập nội dung (bullet points)
    body_shape = slide.placeholders[1]
    
    # Chỉnh lại kích thước khung text nếu có hình ảnh
    if include_image_placeholder:
        body_shape.width = Inches(7.5) # Thu nhỏ lại để nhường chỗ cho ảnh bên phải
        
    tf = body_shape.text_frame
    tf.clear()
    
    for idx, bullet in enumerate(content_bullets):
        p = tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.name = 'Arial'
        p.font.size = Pt(24)
        p.font.color.rgb = RGBColor(51, 51, 51)
        # Tạo space sau mỗi đoạn
        p.space_after = Pt(14)
        
    # Thêm placeholder cho hình ảnh nếu cần
    if include_image_placeholder:
        left = Inches(8.0)
        top = Inches(2.5)
        width = Inches(7.0)
        height = Inches(4.5)
        
        # Vẽ một khung chữ nhật làm placeholder
        rect = slide.shapes.add_shape(
            1, left, top, width, height # 1 is MSO_SHAPE.RECTANGLE
        )
        rect.fill.solid()
        rect.fill.fore_color.rgb = RGBColor(230, 230, 230) # Màu xám nhạt
        
        # Thêm text vào placeholder
        text_frame = rect.text_frame
        text_frame.text = f"[CHÈN HÌNH ẢNH MINH HỌA PHẦN MỀM]\n{image_caption}"
        for paragraph in text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER
            paragraph.font.name = 'Arial'
            paragraph.font.color.rgb = RGBColor(100, 100, 100)
            
    # Thêm note cho diễn giả nếu có
    if note_text:
        notes_slide = slide.notes_slide
        text_frame = notes_slide.notes_text_frame
        text_frame.text = note_text

def main():
    prs = Presentation()
    # Chỉnh tỷ lệ 16:9
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    
    # Slide 1: Tiêu đề
    title_slide_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(title_slide_layout)
    title = slide1.shapes.title
    subtitle = slide1.placeholders[1]
    title.text = "CHUYỂN ĐỔI SỐ\nQUẢN LÝ CÔNG VIỆC & ĐÁNH GIÁ KPI"
    subtitle.text = "Từ Đánh Giá Hình Thức Đến Hiệu Quả Thực Chất"
    
    # Custom font cho Slide 1
    for p in title.text_frame.paragraphs:
        p.font.name = 'Arial'
        p.font.bold = True
        p.font.color.rgb = RGBColor(0, 51, 102)
    for p in subtitle.text_frame.paragraphs:
        p.font.name = 'Arial'
    
    # Slide 2
    create_slide(prs, 
        "Thực Trạng Đánh Giá KPI Hiện Nay",
        [
            "❌ Thủ công & Mất thời gian: Tổng hợp bằng file Word, mất rất nhiều công sức vào cuối mỗi quý/năm.",
            "❌ Mang tính hình thức (Cào bằng): Kết quả gần như ai cũng đạt loại A, không phản ánh đúng năng lực.",
            "❌ Cảm tính: Chấm điểm dựa trên cảm nhận thay vì bám sát danh sách công việc thực tế."
        ],
        note_text="Nhấn mạnh vào nỗi đau mất thời gian và sự thiếu công bằng hiện tại."
    )
    
    # Slide 3
    create_slide(prs,
        "Giải Pháp Phần Mềm Mới",
        [
            "✅ Số hóa toàn diện: Đưa toàn bộ danh mục công việc lên phần mềm, theo dõi liên tục.",
            "✅ Đánh giá dựa trên minh chứng: Dữ liệu thực tế được ghi nhận làm cơ sở vững chắc cho KPI, đảm bảo tính minh bạch và công bằng tuyệt đối."
        ]
    )
    
    # Slide 4
    create_slide(prs,
        "Hệ Thống Phân Quyền Chuyên Biệt",
        [
            "📊 Bảng Tổng Quan (View): Dành cho Sếp xem xét bức tranh toàn cảnh.",
            "📖 Sổ tay Hướng dẫn: Tài liệu tra cứu nội bộ dùng chung.",
            "👤 Quyền Nhân Viên: Tập trung vào Bảng theo dõi tiến độ và Thêm/Cập nhật công việc.",
            "👥 Quyền Quản Lý: Bao gồm duyệt việc, đánh giá KPI, quản lý JD và Cấu hình."
        ],
        include_image_placeholder=True,
        image_caption="(Ảnh 2 menu phân quyền của NV và QL)"
    )
    
    # Slide 5
    create_slide(prs,
        "Quyền Nhân Viên: Tạo Việc Nhanh & Tái Tạo Định Kỳ",
        [
            "• Nhân viên vào Bảng theo dõi tiến độ (có tính năng Sort) để tìm việc.",
            "• Thêm công việc mới cực nhanh.",
            "🌟 Tính năng Tái tạo định kỳ:",
            "  - Tự động hóa các việc lặp lại.",
            "  - Ví dụ: Tháng 9 tạo việc 'Tính lương', tháng 10 chỉ cần bấm 'Tái tạo định kỳ' là hệ thống sinh ra công việc tương tự."
        ],
        include_image_placeholder=True,
        image_caption="(Ảnh tính năng tái tạo định kỳ)"
    )
    
    # Slide 6
    create_slide(prs,
        "Quyền Nhân Viên: Cập Nhật Tiến Độ & Vướng Mắc",
        [
            "✅ Hoàn thành: Cập nhật kết quả kèm theo Link đính kèm minh chứng.",
            "⚠️ Chưa hoàn thành: Bấm vào nút 'Vướng mắc' và điền rõ nguyên nhân.",
            "• Phân loại nguyên nhân: Chọn lý do là Chủ quan hay Khách quan.",
            "• Lưu ý quan trọng: Nếu là lý do Khách quan, hệ thống sẽ tự động gửi thông báo (với các mức đề xuất như 50, 70 điểm...) đến Quản lý để xử lý."
        ],
        include_image_placeholder=True,
        image_caption="(Ảnh giao diện cập nhật vướng mắc)"
    )
    
    # Slide 7
    create_slide(prs,
        "Quyền Quản Lý: Duyệt Việc Khách Quan",
        [
            "• Quản lý xem xét các 'Vướng mắc khách quan' mà nhân viên gửi lên.",
            "• Chấm điểm theo khối lượng công việc đã làm ở các mức độ.",
            "• Tầm quan trọng: Điểm số từ bước duyệt việc khách quan này sẽ được liên kết trực tiếp và ảnh hưởng đến việc chấm điểm KPI cuối cùng."
        ]
    )
    
    # Slide 8
    create_slide(prs,
        "Quy Trình Chấm KPI Cuối Kỳ (TBP & Tổ KPI)",
        [
            "• Trưởng Bộ Phận (TBP) và Tổ KPI tiến hành đánh giá cuối kỳ.",
            "• Phần mềm đề xuất: Hệ thống tự động tính toán điểm số KPI đề xuất dựa trên dữ liệu công việc (gồm cả điểm duyệt khách quan).",
            "• Dò lại: Tổ KPI và TBP chỉ cần dựa trên điểm đề xuất để dò lại, đối chiếu các trường hợp đặc biệt để chốt KPI.",
            "• Xóa bỏ hoàn toàn tình trạng 'A mặc định' cuối năm."
        ]
    )
    
    # Slide 9
    create_slide(prs,
        "Đối Chiếu JD (Mô Tả Công Việc) Bằng AI",
        [
            "• Tổ KPI dùng hệ thống JD để đối chiếu xem nhân viên đăng ký công việc có hợp lý không.",
            "🤖 Tích hợp AI Thông minh: Phần mềm sử dụng AI để 'đọc' và so sánh giữa JD chuẩn với các công việc nhân viên thực tế tạo.",
            "• AI sẽ cảnh báo nếu phát hiện một nhân viên đang làm các công việc sai lệch quá nhiều so với JD ban đầu của họ."
        ]
    )
    
    # Slide 10
    create_slide(prs,
        "Lộ Trình Triển Khai & Áp Dụng",
        [
            "1. 🤝 Họp Ban Lãnh Đạo: Thống nhất chủ trương, lấy ý kiến đóng góp.",
            "2. 📢 Thu thập JD: HCNS yêu cầu các phòng ban rà soát và gửi JD lên hệ thống.",
            "3. 🧪 Thử nghiệm (1-2 tháng): Chạy Pilot thực tế phần mềm, thu thập phản hồi, góp ý từ người dùng và tinh chỉnh.",
            "4. 🚀 Áp dụng chính thức: Chuyển đổi hoàn toàn! Đánh giá KPI 100% trên phần mềm."
        ]
    )
    
    prs.save('Ban_Thuyet_Trinh_Phan_Mem_KPI.pptx')
    print("Presentation saved as Ban_Thuyet_Trinh_Phan_Mem_KPI.pptx")

if __name__ == '__main__':
    main()
