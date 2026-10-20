# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_full_document_v3():
    doc = docx.Document()
    
    # Setup page margins
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(4)

    def set_cell_bg(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_paddings(cell, top=120, bottom=120, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'''
            <w:tcMar {nsdecls("w")}>
                <w:top w:w="{top}" w:type="dxa"/>
                <w:bottom w:w="{bottom}" w:type="dxa"/>
                <w:left w:w="{left}" w:type="dxa"/>
                <w:right w:w="{right}" w:type="dxa"/>
            </w:tcMar>
        ''')
        tcPr.append(tcMar)

    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(8)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.bold = True
        r.font.name = 'Arial'
        r.font.size = Pt(15)
        r.font.color.rgb = RGBColor(10, 37, 64) # Navy Blue
        pPr = h._p.get_or_add_pPr()
        pBdr = parse_xml(f'''
            <w:pBdr {nsdecls("w")}>
                <w:bottom w:val="single" w:sz="16" w:space="6" w:color="0052CC"/>
            </w:pBdr>
        ''')
        pPr.append(pBdr)
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(13)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.bold = True
        r.font.name = 'Arial'
        r.font.size = Pt(12.5)
        r.font.color.rgb = RGBColor(0, 82, 204) # Royal Blue
        return h

    def add_callout(title, bullets, border_color="0052CC", bg_color="F0F7FF"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.8)
        set_cell_bg(cell, bg_color)
        set_cell_paddings(cell, 120, 120, 160, 160)
        
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        tcPr.append(tcBorders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        r_t = p.add_run(f"📌 {title}")
        r_t.bold = True
        r_t.font.name = 'Arial'
        r_t.font.size = Pt(10.5)
        r_t.font.color.rgb = RGBColor(10, 37, 64)
        
        for b in bullets:
            p_b = cell.add_paragraph()
            p_b.paragraph_format.space_before = Pt(1)
            p_b.paragraph_format.space_after = Pt(2)
            p_b.paragraph_format.line_spacing = 1.15
            r_b = p_b.add_run(b)
            r_b.font.name = 'Arial'
            r_b.font.size = Pt(9.5)
            r_b.font.color.rgb = RGBColor(51, 65, 85)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ================= COVER / TITLE =================
    p_org = doc.add_paragraph()
    p_org.paragraph_format.space_after = Pt(2)
    p_org.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_org = p_org.add_run("TỔNG CÔNG TY DMT GROUP — BAN HÀNH CHÍNH NHÂN SỰ\n")
    r_org.bold = True
    r_org.font.size = Pt(11)
    r_org.font.color.rgb = RGBColor(100, 116, 139)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(6)
    p_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("KỊCH BẢN THUYẾT TRÌNH & HƯỚNG DẪN VẬN HÀNH HỆ THỐNG KPI\n")
    r_title.bold = True
    r_title.font.size = Pt(17)
    r_title.font.color.rgb = RGBColor(10, 37, 64)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(14)
    p_sub.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("DÀNH CHO CÁC ANH CHỊ TRƯỞNG BAN VÀ TOÀN THỂ CÁN BỘ NHÂN VIÊN\n")
    r_sub.bold = True
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = RGBColor(0, 82, 204)
    r_sub2 = p_sub.add_run("Thực hành thao tác • Nắm vững 5 nguyên tắc vận hành • Đối chiếu cách làm cũ & mới")
    r_sub2.italic = True
    r_sub2.font.size = Pt(10.5)
    r_sub2.font.color.rgb = RGBColor(71, 85, 105)

    # Info box
    info_table = doc.add_table(rows=2, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False
    for r in info_table.rows:
        for c in r.cells:
            c.width = Inches(3.4)
            set_cell_bg(c, "F8FAFC")
            set_cell_paddings(c, 80, 80, 100, 100)
    
    info_table.cell(0, 0).paragraphs[0].text = "🎯 Thành phần tham dự: Các anh/chị Trưởng Ban & Toàn thể CBNV"
    info_table.cell(0, 1).paragraphs[0].text = "⏱️ Thời lượng: 40 - 45 phút (Thuyết trình + Demo + Q&A)"
    info_table.cell(1, 0).paragraphs[0].text = "🎙️ Đơn vị tổ chức & Trình bày: Ban Hành chính Nhân sự (HCNS)"
    info_table.cell(1, 1).paragraphs[0].text = "📢 Tinh thần cốt lõi: Tập trung thực hành — Minh bạch — Đối chiếu thực tiễn"
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ================= PHẦN 1 =================
    add_h1("PHẦN 1. KHAI MẠC VÀ GIỚI THIỆU MỤC TIÊU (5 PHÚT)")
    
    add_callout("MỤC TIÊU CỐT LÕI CỦA PHẦN 1", [
        "1. Tạo bầu không khí cởi mở, trang trọng nhưng gần gũi, xóa bỏ định kiến e ngại về phần mềm mới.",
        "2. Tuyên bố rõ phương châm: KHÔNG đi sâu lý thuyết suông -> Tập trung 100% vào thao tác thực chiến và xử lý tình huống.",
        "3. Nhấn mạnh tinh thần Quý 4/2026: Giai đoạn chạy thử nghiệm (Pilot) để làm quen tay, yên tâm trải nghiệm, HCNS đồng hành sát sao."
    ], border_color="0052CC", bg_color="F0F7FF")

    t1 = doc.add_table(rows=3, cols=1)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    t1.autofit = False
    
    c0 = t1.cell(0, 0)
    c0.width = Inches(6.8)
    set_cell_bg(c0, "F1F5F9")
    set_cell_paddings(c0, 80, 80, 120, 120)
    p0 = c0.paragraphs[0]
    p0.add_run("📺 NỘI DUNG TRÌNH CHIẾU SLIDE 1 (KHAI MẠC & MỤC TIÊU):").bold = True
    bullets_s1 = [
        "• CHỦ ĐỀ: HƯỚNG DẪN SỬ DỤNG HỆ THỐNG QUẢN LÝ TIẾN ĐỘ & ĐÁNH GIÁ KPI (DMT GROUP)",
        "• 3 TRỌNG TÂM BUỔI LÀM VIỆC:",
        "   1. Hiểu đúng vai trò và trách nhiệm của từng vị trí trên hệ thống.",
        "   2. Hướng dẫn thao tác trực tiếp: Tạo việc, Cập nhật tiến độ, Duyệt việc & Nghiệm thu Chỉ tiêu.",
        "   3. Thống nhất cách xử lý các tình huống thực tế thường gặp.",
        "• TRẢI NGHIỆM CÁ NHÂN HÓA: Chọn đúng Công ty -> đúng Tên -> đúng Vai trò (Giao diện tinh gọn).",
        "• TINH THẦN QUÝ 4/2026: Giai đoạn chạy thử (Pilot) — Trưởng ban làm quen chia nhỏ mục tiêu, nhân viên tạo thói quen cập nhật — Ban HCNS hỗ trợ 24/7!"
    ]
    for b in bullets_s1:
        pb = c0.add_paragraph()
        pb.paragraph_format.space_before = Pt(1)
        pb.paragraph_format.space_after = Pt(2)
        pb.add_run(b).font.size = Pt(9.5)

    c1 = t1.cell(1, 0)
    c1.width = Inches(6.8)
    set_cell_bg(c1, "FFFFFF")
    set_cell_paddings(c1, 100, 100, 120, 120)
    p1 = c1.paragraphs[0]
    r1 = p1.add_run("🎙️ LỜI TRÌNH BÀY DIỄN GIẢ (SPEAKER SCRIPT NÂNG CẤP):")
    r1.bold = True
    r1.font.color.rgb = RGBColor(0, 82, 204)

    script_p1 = [
        "\"Kính chào các anh chị Trưởng Ban cùng toàn thể anh chị em đại gia đình DMT Group!",
        "Thay mặt Ban Hành chính Nhân sự, em xin gửi lời cảm ơn chân thành đến tất cả các anh chị em đã sắp xếp thời gian có mặt đông đủ trong buổi hướng dẫn sử dụng Hệ thống Quản lý tiến độ và Đánh giá KPI ngày hôm nay.",
        "Trước đó, toàn bộ lộ trình triển khai và quy chế vận hành đã được báo cáo chi tiết lên Ban Lãnh đạo và đạt được sự thống nhất cao. Các tài liệu hướng dẫn nghiệp vụ cũng đã được Ban HCNS gửi trước để các anh chị nghiên cứu. Vì vậy, trong buổi làm việc hôm nay, chúng ta sẽ [nhấn mạnh] hoàn toàn KHÔNG đi sâu trình bày lại lý thuyết suông, mà sẽ tập trung 100% vào THAO TÁC THỰC TẾ và CÁCH XỬ LÝ NHỮNG TÌNH HUỐNG PHÁT SINH trong công việc hàng ngày.",
        "Cụ thể, buổi hướng dẫn hôm nay gồm 3 nội dung trọng tâm:",
        "• Thứ nhất: Giúp mọi người hiểu thật rõ vai trò và trách nhiệm cụ thể của mình trên hệ thống.",
        "• Thứ hai: Cầm tay chỉ việc, thao tác trực tiếp trên màn hình để anh chị em biết cách tạo việc, cập nhật tiến độ hàng ngày; các cấp Quản lý biết cách duyệt việc và nghiệm thu Chỉ tiêu theo đúng thẩm quyền được phân quyền.",
        "• Thứ ba: Thống nhất nguyên tắc xử lý nhanh gọn các tình huống thường gặp khi áp dụng vào thực tế.",
        "Hệ thống này được triển khai dùng chung cho tất cả các đơn vị đã tích hợp trong tập đoàn. Giao diện được thiết kế cá nhân hóa và cực kỳ thông minh: Khi đăng nhập, mọi người chỉ cần chọn đúng Tên công ty, đúng Họ tên của mình và đúng Vai trò được cấp. Ngay lập tức, hệ thống sẽ tự động lọc và hiển thị chính xác các dữ liệu liên quan, tuyệt đối không làm rối mắt hay gây nhầm lẫn.",
        "Và cuối cùng, em xin được chia sẻ một thông điệp giải tỏa tâm lý quan trọng: Trong Quý 4 năm nay, chúng ta mới chỉ chính thức bước vào GIAI ĐOẠN CHẠY THỬ (PILOT)! Đây là khoảng thời gian để các anh chị Trưởng ban làm quen với kỹ năng chia nhỏ mục tiêu trên phần mềm, và để anh chị em nhân viên tập thói quen cập nhật tiến độ công việc liên tục thay vì làm thủ công.",
        "Vì là một công cụ mới, chắc chắn trong thời gian đầu sẽ có lúc chúng ta thao tác nhầm hoặc còn bỡ ngỡ. Các anh chị và các bạn cứ hoàn toàn yên tâm dùng thử thật thoải mái! Bất kỳ khi nào gặp vướng mắc, mọi người chỉ cần nhắn lại cho Ban HCNS. Đội ngũ kỹ thuật và HCNS luôn sẵn sàng hỗ trợ tức thì và tiếp thu ý kiến của mọi người để hoàn thiện hệ thống ngày càng tiện ích hơn ạ!\""
    ]
    for sp in script_p1:
        psp = c1.add_paragraph()
        psp.paragraph_format.space_before = Pt(3)
        psp.paragraph_format.space_after = Pt(4)
        psp.paragraph_format.line_spacing = 1.2
        psp.add_run(sp).font.size = Pt(9.5)

    c2 = t1.cell(2, 0)
    c2.width = Inches(6.8)
    set_cell_bg(c2, "FEF3C7")
    set_cell_paddings(c2, 60, 60, 120, 120)
    p2 = c2.paragraphs[0]
    p2.add_run("💡 Mẹo trình bày sân khấu: Nở nụ cười thân thiện, ánh mắt mở đầu hướng về phía các Trưởng Ban để thể hiện sự tôn trọng, sau đó quét rộng bao quát toàn thể hội trường. Ở câu cuối về cam kết Quý 4 chạy thử, hãy nói với giọng điệu trấn an, chân thành để giải tỏa 100% tâm lý lo sợ bấm nhầm bị phạt của nhân viên.").font.size = Pt(9)
    p2.runs[0].font.italic = True
    p2.runs[0].font.color.rgb = RGBColor(146, 64, 14)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ================= PHẦN 2: CÓ SO SÁNH VỚI CÁCH LÀM CŨ =================
    add_h1("PHẦN 2. NGUYÊN TẮC VẬN HÀNH CẦN HIỂU CHUNG (7 PHÚT)")
    
    add_callout("ĐIỂM NÂNG CẤP ĐẶC BIỆT CỦA PHẦN 2", [
        "Nêu bật sự ĐỐI CHIẾU giữa CÁCH LÀM CŨ (file Word/Excel, dồn toa, duyệt miệng) với NGUYÊN TẮC MỚI trên phần mềm.",
        "Giúp Trưởng ban và nhân viên hiểu rõ bản chất 'Tại sao phải làm như vậy' thay vì cảm thấy bị gò bó bởi quy trình mới."
    ], border_color="0052CC", bg_color="F0F7FF")

    add_h2("Bảng đối chiếu nhanh: Cách làm cũ vs 5 Nguyên tắc mới")
    tbl_cmp = doc.add_table(rows=6, cols=3)
    tbl_cmp.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_cmp.autofit = False
    
    headers = ["Nội dung", "Cách làm cũ (Thủ công / Truyền thống)", "Nguyên tắc mới trên phần mềm"]
    widths = [Inches(1.5), Inches(2.6), Inches(2.7)]
    for idx, h_text in enumerate(headers):
        c = tbl_cmp.cell(0, idx)
        c.width = widths[idx]
        set_cell_bg(c, "0A2540")
        set_cell_paddings(c, 100, 100, 100, 100)
        p = c.paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

    cmp_rows = [
        ("1. Nguồn gốc công việc", "Nhân viên tự nghĩ và liệt kê việc rời rạc; nhiều việc không rõ phục vụ mục tiêu nào của phòng.", "Liên kết trục dọc: Dự án ➔ Chỉ tiêu Ban ➔ Công việc cá nhân. Làm việc nào rõ đích đến việc đó."),
        ("2. Trách nhiệm tạo & làm việc", "Nhân viên coi báo cáo là làm cho sếp xem; sếp giao việc miệng qua Zalo/email dễ quên, không theo dõi việc của chính sếp.", "Trách nhiệm 2 chiều: Cả sếp và nhân viên đều làm việc; sếp phân rã chỉ tiêu và tự cập nhật việc chuyên môn của mình."),
        ("3. Thời điểm cập nhật", "Thói quen 'dồn toa' đến cuối tháng mới ngồi nhớ lại để gõ báo cáo; vướng mắc để dồn cuối kỳ mới vỡ lở.", "Cập nhật Real-time ngay khi làm; có nguy cơ trễ hoặc vướng mắc phải báo cáo sớm để sếp hỗ trợ tháo gỡ."),
        ("4. Minh chứng & Nghiệm thu", "Khai báo 'Đã xong' chung chung trên giấy tờ; duyệt theo cảm tính hoặc tin tưởng miệng, thiếu cơ sở đối chứng.", "Báo hoàn thành bắt buộc đính kèm File/Link minh chứng thực tế; thước đo nghiệm thu là kết quả thực tế ngoài đời."),
        ("5. Cấp độ đánh giá", "Nhầm lẫn: Nhân viên xong hết việc cá nhân là nghĩ Chỉ tiêu của Ban xong; dẫn đến tỷ lệ ảo dù dự án chung bị chậm.", "Phân định 2 cấp độ rõ ràng: Việc cá nhân xong KHÔNG mặc nhiên là Chỉ tiêu của Ban xong; sếp phải nghiệm thu Master View.")
    ]

    for row_idx, data in enumerate(cmp_rows, start=1):
        bg = "FFFFFF" if row_idx % 2 == 1 else "F8FAFC"
        for col_idx in range(3):
            c = tbl_cmp.cell(row_idx, col_idx)
            c.width = widths[col_idx]
            set_cell_bg(c, bg)
            set_cell_paddings(c, 80, 80, 100, 100)
            p = c.paragraphs[0]
            r = p.add_run(data[col_idx])
            r.font.size = Pt(9.5)
            if col_idx == 0:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    t2 = doc.add_table(rows=3, cols=1)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t2.autofit = False
    
    c0 = t2.cell(0, 0)
    c0.width = Inches(6.8)
    set_cell_bg(c0, "F1F5F9")
    set_cell_paddings(c0, 80, 80, 120, 120)
    p0 = c0.paragraphs[0]
    p0.add_run("📺 NỘI DUNG TRÌNH CHIẾU SLIDE 2 & 3 (5 NGUYÊN TẮC VẬN HÀNH & ĐỐI CHIẾU CÁCH LÀM CŨ):").bold = True
    bullets_s2 = [
        "• 5 NGUYÊN TẮC VẬN HÀNH CỐT LÕI (SO SÁNH CŨ vs MỚI):",
        "  1. XUẤT PHÁT TỪ KẾ HOẠCH NĂM & CHỈ TIÊU BAN (Khác cách cũ: Không tự nghĩ việc rời rạc ➔ Đi theo 3 bước: Dự án ➔ Chỉ tiêu ➔ Tên việc).",
        "  2. TRÁCH NHIỆM HAI CHIỀU (Khác cách cũ: Không chỉ nhân viên báo cáo sếp ➔ Cả sếp và nhân viên đều tạo & làm việc của mình trên máy).",
        "  3. CẬP NHẬT TIẾN ĐỘ REAL-TIME (Khác cách cũ: Chấm dứt 'dồn toa' cuối tháng ➔ Báo cáo tiến độ liên tục, báo sớm vướng mắc để tháo gỡ).",
        "  4. HOÀN THÀNH PHẢI CÓ MINH CHỨNG THỰC TẾ (Khác cách cũ: Không khai 'Đã xong' bằng miệng ➔ Bắt buộc đính kèm File/Link minh chứng).",
        "  5. PHÂN BIỆT RÕ 2 CẤP ĐỘ ĐÁNH GIÁ (Khác cách cũ: Không ngộ nhận việc cá nhân xong là Ban xong ➔ Quản lý phải trực tiếp nghiệm thu Chỉ tiêu)."
    ]
    for b in bullets_s2:
        pb = c0.add_paragraph()
        pb.paragraph_format.space_before = Pt(1)
        pb.paragraph_format.space_after = Pt(2)
        pb.add_run(b).font.size = Pt(9.5)

    c1 = t2.cell(1, 0)
    c1.width = Inches(6.8)
    set_cell_bg(c1, "FFFFFF")
    set_cell_paddings(c1, 100, 100, 120, 120)
    p1 = c1.paragraphs[0]
    r1 = p1.add_run("🎙️ LỜI TRÌNH BÀY DIỄN GIẢ (SPEAKER SCRIPT CÓ LỒNG GHÉP SO SÁNH VỚI CÁCH LÀM CŨ):")
    r1.bold = True
    r1.font.color.rgb = RGBColor(0, 82, 204)

    script_p2 = [
        "\"Thưa các anh chị, trước khi bước vào thực hành các thao tác cụ thể trên phần mềm, chúng ta cần cùng nhau thống nhất 5 NGUYÊN TẮC CHUNG. Để mọi người dễ hình dung lý do tại sao hệ thống lại thiết kế như vậy, em xin phép vừa nêu nguyên tắc, vừa đối chiếu nhanh với cách làm truyền thống trước đây của chúng ta:",
        "NGUYÊN TẮC THỨ NHẤT: Công việc của từng cá nhân bắt buộc phải xuất phát từ Kế hoạch năm và Chỉ tiêu chung của Phòng/Ban.",
        "• Trước đây với cách làm cũ: Mỗi cá nhân thường tự ngồi nghĩ ra danh sách việc hàng tháng của mình trên file Word hoặc Excel một cách khá rời rạc. Hậu quả là có những việc làm rất mất thời gian nhưng nhìn lại thì không rõ nó đang phục vụ cho mục tiêu trọng tâm nào của phòng hay của công ty.",
        "• Còn trên phần mềm mới: Hệ thống yêu cầu chúng ta tuân thủ đúng trình tự 3 bước: Chọn Dự án ➔ Chọn Chỉ tiêu ➔ rồi mới nhập Tên công việc cụ thể. Việc liên kết chặt chẽ này đảm bảo mỗi việc anh chị em làm ra đều có căn cứ rõ ràng, không bị làm việc lan man, và giúp Ban Lãnh đạo nhìn thấy rõ từng nỗ lực hàng ngày của các bạn đang đóng góp trực tiếp vào Chỉ tiêu quan trọng nào của tổ chức.",
        "NGUYÊN TẮC THỨ HAI: Trách nhiệm hai chiều — Cả Nhân viên và Quản lý đều phải tạo việc và thực hiện công việc trên hệ thống.",
        "• Trước đây với cách làm cũ: Nhiều người có tâm lý coi việc làm báo cáo là 'làm để nộp cho sếp xem'. Còn các sếp thì chủ yếu giao việc bằng miệng, qua điện thoại hoặc tin nhắn Zalo, rất dễ thất lạc và bản thân Quản lý cũng không có chỗ nào để ghi nhận các đầu việc điều hành thực tế của chính mình.",
        "• Còn trên phần mềm mới: Trách nhiệm là hai chiều sòng phẳng. Nhân viên tiếp nhận việc được giao và chủ động tạo các việc thường nhật thuộc trách nhiệm của mình. Còn các anh chị Quản lý [nhìn về phía các Trưởng Ban và mỉm cười] cũng không chỉ ngồi theo dõi cấp dưới, mà chính Quản lý cũng phải tạo và cập nhật các công việc chuyên môn của bản thân; đồng thời chịu trách nhiệm 'phân rã' Chỉ tiêu lớn thành các đầu việc nhỏ, phân công đúng người, theo sát tiến độ và nghiệm thu kết quả.",
        "NGUYÊN TẮC THỨ BA: Cập nhật tiến độ phải được thực hiện theo thời gian thực (Real-time), tuyệt đối không để dồn đến cuối tháng mới cập nhật.",
        "• Trước đây với cách làm cũ: Căn bệnh kinh niên của chúng ta là thói quen 'dồn toa'. Suốt cả tháng không ghi chép gì, đến sát ngày chốt sổ mới cuống cuồng lục lại email, vắt óc nhớ lại xem mình đã làm gì. Nguy hiểm hơn, những công việc bị tắc nghẽn hay vướng mắc cũng bị ỉm đi, đợi đến cuối kỳ mới báo thì sự việc đã rồi, cấp trên không kịp ra tay hỗ trợ.",
        "• Còn trên phần mềm mới: Tiến độ được cập nhật liên tục ngay trong quá trình làm việc. Đặc biệt, khi nhận thấy một công việc có nguy cơ trễ hạn hoặc xuất hiện vướng mắc khách quan (như chờ phản hồi đối tác, vướng thủ tục sở ban ngành, thời tiết...), nhân viên cần báo cáo sớm ngay trên hệ thống: nêu rõ tình trạng, nguyên nhân và hướng đề xuất. Báo cáo sớm trên phần mềm mới không phải để bị phạt, mà là để Quản lý kịp thời nhảy vào hỗ trợ tháo gỡ khó khăn cho các bạn!",
        "NGUYÊN TẮC THỨ TƯ: Khi báo 'Hoàn thành', bắt buộc phải cập nhật kết quả và đính kèm minh chứng thực tế.",
        "• Trước đây với cách làm cũ: Chúng ta chỉ đánh dấu 'Đã xong' trên một dòng chữ trong bảng biểu. Việc đánh giá và duyệt chủ yếu dựa trên niềm tin hoặc cảm nhận chủ quan, không có gì để đối chiếu chất lượng.",
        "• Còn trên phần mềm mới: Hệ thống hướng tới hiệu quả thực chất. Khi báo Hoàn thành, nhân viên cần đính kèm file tài liệu, đường link văn bản hoặc hình ảnh nghiệm thu thực tế phù hợp. Về phía Quản lý, các anh chị có trách nhiệm kiểm tra nội dung và minh chứng đính kèm trước khi bấm nút Duyệt việc. Trạng thái 'Hoàn thành' trên máy chỉ là thông tin hiển thị; việc nghiệm thu công việc của chúng ta bắt buộc phải căn cứ vững chắc vào kết quả thực tế đạt được.",
        "VÀ NGUYÊN TẮC THỨ NĂM [chậm lại, nhấn mạnh] — Điểm khác biệt lớn nhất so với cách làm cũ: Phân biệt rõ hai cấp độ đánh giá trên hệ thống.",
        "• Trước đây với cách làm cũ: Chúng ta rất dễ rơi vào sự ngộ nhận và cào bằng. Cứ thấy nhân viên trong phòng báo cáo hoàn thành hết việc cá nhân là mặc nhiên kết luận phòng mình đã hoàn thành xuất sắc nhiệm vụ tháng đó. Nhưng thực tế, Kế hoạch lớn của dự án hay mục tiêu chung của công ty lại bị trễ hạn, tạo ra một sự 'nghịch lý điểm cao nhưng việc chung không chạy'.",
        "• Còn trên phần mềm mới: Hệ thống tách bạch rạch ròi 2 cấp độ đánh giá:",
        "   - Cấp độ 1 là KẾT QUẢ CÔNG VIỆC CÁ NHÂN: Đánh giá xem từng cá nhân đã hoàn thành các đầu việc được giao như thế nào.",
        "   - Cấp độ 2 là KẾT QUẢ CHỈ TIÊU CỦA PHÒNG/BAN: Tổng hợp tất cả các công việc cá nhân có thực sự giúp Ban đạt được mục tiêu Kế hoạch năm/quý đề ra hay chưa.",
        "Chính vì vậy, anh chị em cần hiểu rõ: Việc tất cả nhân viên trong phòng đều làm xong 100% công việc của mình KHÔNG MẶC NHIÊN đồng nghĩa với việc Chỉ tiêu chung của Ban đã hoàn thành! Dù thanh tiến độ có chạy đủ 100%, Chỉ tiêu vẫn sẽ không tự động đóng lại. Trưởng ban bắt buộc phải là người trực tiếp thẩm định chất lượng tổng thể cuối cùng và đích thân bấm nút xác nhận 'Nghiệm thu Chỉ tiêu' trên Master View. Điều này đảm bảo toàn công ty chúng ta hướng tới hiệu quả thực chất, chứ không chạy theo tỷ lệ phần trăm ảo trên phần mềm ạ!\""
    ]
    for sp in script_p2:
        psp = c1.add_paragraph()
        psp.paragraph_format.space_before = Pt(3)
        psp.paragraph_format.space_after = Pt(4)
        psp.paragraph_format.line_spacing = 1.2
        psp.add_run(sp).font.size = Pt(9.5)

    c2 = t2.cell(2, 0)
    c2.width = Inches(6.8)
    set_cell_bg(c2, "FEF3C7")
    set_cell_paddings(c2, 60, 60, 120, 120)
    p2 = c2.paragraphs[0]
    p2.add_run("💡 Mẹo trình bày sân khấu: Khi nói đến 'Trước đây cách làm cũ', hãy dùng giọng điệu đồng cảm, hơi vui vẻ một chút để thừa nhận những thói quen cũ mà ai cũng từng mắc phải. Khi chuyển sang 'Còn trên phần mềm mới', hãy nói với giọng dứt khoát, tự tin, nhấn mạnh lợi ích minh bạch và công bằng cho cả Sếp và Nhân viên.").font.size = Pt(9)
    p2.runs[0].font.italic = True
    p2.runs[0].font.color.rgb = RGBColor(146, 64, 14)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)
    
    # Save file with fallback if file is locked
    primary_file = "Kich_Ban_Thuyet_Trinh_Toan_The_Nhan_Vien.docx"
    backup_file = "Kich_Ban_Thuyet_Trinh_Toan_The_Nhan_Vien_Moi.docx"
    try:
        doc.save(primary_file)
        print(f"File saved successfully to: {primary_file}")
    except PermissionError:
        doc.save(backup_file)
        print(f"Primary file locked by Word. Saved successfully to: {backup_file}")

if __name__ == "__main__":
    create_full_document_v3()
