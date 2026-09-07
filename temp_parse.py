import docx
import sys

def parse_docx(filepath, outpath):
    doc = docx.Document(filepath)
    with open(outpath, 'w', encoding='utf-8') as f:
        f.write(f"--- Paragraphs ({len(doc.paragraphs)}) ---\n")
        for i, p in enumerate(doc.paragraphs):
            text = p.text.strip()
            if text:
                f.write(f"[{i}] {text}\n")
        
        f.write(f"\n--- Tables ({len(doc.tables)}) ---\n")
        for i, t in enumerate(doc.tables):
            f.write(f"\nTable {i} ({len(t.rows)} rows):\n")
            for row in t.rows:
                cells = [cell.text.replace('\n', ' ').strip() for cell in row.cells]
                f.write(" | ".join(cells) + "\n")

if __name__ == '__main__':
    filepath = r'c:\Users\Admin\Desktop\AG\Theodoitiendo\INPUT\HÀNH CHÍNH NHÂN SỰ\JD Le Thi Tu Uyen ( Final ).docx'
    outpath = r'c:\Users\Admin\Desktop\AG\Theodoitiendo\out_docx.txt'
    parse_docx(filepath, outpath)
    print("Done")
