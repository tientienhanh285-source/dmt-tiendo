import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def main():
    doc = Document()
    
    # Tiêu đề
    title = doc.add_heading('LỘ TRÌNH ĐÁNH GIÁ VÀ QUY CHẾ VẬN HÀNH HỆ THỐNG KPI - DMT GROUP', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # PHẦN I
    doc.add_heading('PHẦN I: LỘ TRÌNH ĐÁNH GIÁ VÀ VẬN HÀNH HỆ THỐNG', level=1)
    
    doc.add_paragraph('Cách vận hành của hệ thống được tổ chức theo một trục dọc xuyên suốt, liên kết từ Mục tiêu của Công ty → Chỉ tiêu của Phòng/Ban → Công việc của từng nhân viên → Kết quả thực hiện, bảo đảm công việc hằng tháng được gắn với kế hoạch chung của Công ty.')
    
    doc.add_heading('Tầng 1: Đầu năm & Đầu Quý: Thiết lập, rà soát và điều chỉnh Mục tiêu – Cấp Công ty và Phòng/Ban', level=2)
    p1 = doc.add_paragraph(style='List Bullet')
    p1.add_run('Đầu năm – Thiết lập Kế hoạch Tổng: ').bold = True
    p1.add_run('Ban Lãnh đạo (BLĐ) phê duyệt Bảng Kế hoạch Năm (Master View), trong đó xác định các Chỉ tiêu cụ thể cho từng Phòng/Ban (KHĐT, Kỹ thuật, ĐBGT, v.v.). Các Chỉ tiêu này là kim chỉ nam và đích đến của toàn hệ thống. Đối với các Chỉ tiêu liên quan đến dự án, phải xác định rõ mốc thời gian và thời hạn hoàn thành để làm căn cứ theo dõi và đánh giá (thường là theo Tháng hoặc Quý).')
    p2 = doc.add_paragraph(style='List Bullet')
    p2.add_run('Đầu mỗi Quý – Rà soát và điều chỉnh: ').bold = True
    p2.add_run('Tổ chức họp KPI cấp Công ty/Phòng/Ban để tổng kết kết quả Quý trước, đánh giá tình hình thực hiện các Chỉ tiêu, xác định Chỉ tiêu đã hoàn thành, chưa hoàn thành và nguyên nhân. Trên cơ sở kết quả rà soát và tình hình thực tế, Công ty/Phòng/Ban triển khai Kế hoạch Quý tiếp theo; trường hợp cần thiết có thể giữ nguyên, điều chỉnh hoặc cập nhật Chỉ tiêu và tiến độ thực hiện theo thẩm quyền.')
    
    doc.add_heading('Tầng 2: Đầu mỗi tháng: Phân rã Chỉ tiêu thành Công việc – Cấp Nhân viên', level=2)
    doc.add_paragraph('Quản lý: Căn cứ Kế hoạch đang được theo dõi trên Master View, vào đầu mỗi tháng, Quản lý phân rã các Chỉ tiêu thành các đầu việc cụ thể và giao cho từng nhân viên thực hiện.', style='List Bullet')
    p3 = doc.add_paragraph(style='List Bullet')
    p3.add_run('Nhân viên: Quy trình tạo việc 3 bước: ').bold = True
    p3.add_run('Để bảo đảm công việc được giao có căn cứ và phục vụ mục tiêu chung, hệ thống yêu cầu công việc phải được gắn theo trình tự: Chọn Dự án → Chọn Chỉ tiêu → Nhập tên công việc. Đối với công việc phát sinh ngoài kế hoạch, công việc phải được xác định rõ nội dung, căn cứ phát sinh và được Trưởng Ban/Quản lý xác nhận trước khi đưa vào đánh giá KPI. (Lưu ý: Các công việc phát sinh này sẽ được nhân viên gán vào một Chỉ tiêu chung mang tên "Công việc vận hành, phát sinh theo chỉ đạo" do Ban Lãnh đạo khởi tạo từ đầu năm).')
    p4 = doc.add_paragraph(style='List Bullet')
    p4.add_run('Công việc định kỳ: ').bold = True
    p4.add_run('Đối với các công việc có tính chất lặp lại (ví dụ: tính công, tính lương, báo cáo định kỳ...), Nhân sự/Quản lý có thể tạo đầu mục công việc định kỳ. Sang tháng tiếp theo, hệ thống cho phép sử dụng chức năng "Tái tạo định kỳ" để sao chép công việc sang kỳ mới, giảm thời gian nhập liệu.')
    
    doc.add_heading('Tầng 3: Cuối tháng & Cuối Quý: Kiểm soát tiến độ, nghiệm thu và đánh giá kết quả', level=2)
    p5 = doc.add_paragraph(style='List Bullet')
    p5.add_run('Cuối mỗi tháng: ').bold = True
    p5.add_run('Phần mềm tự động tổng hợp kết quả thực hiện công việc của từng nhân viên trên cơ sở các công việc đã được giao, hoàn thành và được Quản lý duyệt việc (Approve Task); từ đó làm căn cứ tính điểm và xếp loại KPI tháng theo quy định.')
    p6 = doc.add_paragraph(style='List Bullet')
    p6.add_run('Cuối mỗi Quý: ').bold = True
    p6.add_run('Hệ thống tổng hợp dữ liệu KPI và tiến độ thực hiện công việc để đánh giá KPI trong 3 tháng của quý đó. Dựa trên Kế hoạch theo quý để đánh giá các chỉ tiêu đạt/chưa đạt, từ đó truy ngược lại những chỉ tiêu chưa hoàn thành là do cá nhân nào chịu trách nhiệm, làm cơ sở trừ điểm và có thể thiết lập lại mục tiêu Quý sau bằng việc đưa các chỉ tiêu chưa hoàn thành của Quý trước qua sắp xếp, hoặc giữ nguyên tùy tình hình hoạt động, nhằm phục vụ cuộc họp rà soát đầu Quý tiếp theo.')
    p7 = doc.add_paragraph(style='List Bullet')
    p7.add_run('Lợi ích cốt lõi của Lộ trình (Giảm tải áp lực cuối năm): ').bold = True
    p7.add_run('Việc đánh giá, chốt điểm và rà soát chỉ tiêu được làm dứt điểm và cuốn chiếu theo từng Quý. Do đó, hệ thống sẽ tự động cộng dồn và tính toán điểm số cuối năm. Thay vì phải "bơi" trong núi công việc, truy vết lại dữ liệu từ đầu năm một cách quá tải như cách làm cũ, buổi họp cuối năm giờ đây chỉ mang tính chất tổng quan, nhìn lại bức tranh toàn cảnh để định hướng chiến lược cho năm sau.')
    
    doc.add_paragraph('Việc đánh giá được thực hiện ở 02 cấp độ:')
    doc.add_paragraph('Cấp độ 1 – Kết quả công việc: Các công việc được giao cho nhân viên đã được thực hiện và nghiệm thu như thế nào.', style='List Bullet')
    doc.add_paragraph('Cấp độ 2 – Kết quả Chỉ tiêu: Kết quả cộng gộp của các công việc có thực sự bảo đảm hoàn thành Chỉ tiêu của Phòng/Ban hay chưa. Việc hoàn thành công việc của cá nhân không mặc nhiên được xem là hoàn thành Chỉ tiêu của Ban. Chỉ tiêu chỉ được xác nhận hoàn thành sau khi Quản lý đánh giá kết quả tổng thể và xác nhận trên hệ thống.', style='List Bullet')
    
    doc.add_heading('CƠ CHẾ TRÁCH NHIỆM CHÉO', level=2)
    doc.add_paragraph('Kết quả thực hiện Chỉ tiêu của Phòng/Ban là một trong những căn cứ bắt buộc khi đánh giá trách nhiệm của Trưởng Ban.', style='List Bullet')
    doc.add_paragraph('Trường hợp Chỉ tiêu Quý không hoàn thành, Trưởng Ban/Quản lý phải xác định nguyên nhân và trách nhiệm liên quan. Các cá nhân có liên can về việc không hoàn thành chỉ tiêu và Quản lý (việc phân rã Chỉ tiêu, giao việc, theo dõi hoặc điều phối chưa phù hợp), kết quả đánh giá KPI của Trưởng Ban và các cá nhân này sẽ được xem xét điều chỉnh tương ứng.', style='List Bullet')
    doc.add_paragraph('Tức là dù trong 3 tháng của Quý các cá nhân được A, nhưng chỉ tiêu Quý không hoàn thành thì sẽ xem xét điểm xếp loại của quản lý và các cá nhân đó, tiến hành trừ điểm theo quy chế.', style='List Bullet')
    doc.add_paragraph('Trường hợp Chỉ tiêu không hoàn thành do nguyên nhân khách quan hoặc nguyên nhân nằm ngoài phạm vi kiểm soát của Phòng/Ban, phải được ghi nhận, giải trình và xác nhận trên hệ thống.', style='List Bullet')
    doc.add_paragraph('Cơ chế này nhằm bảo đảm toàn bộ hệ thống cùng hướng đến một mục tiêu chung: hoàn thành Kế hoạch của Công ty, thay vì chỉ tập trung vào việc hoàn thành riêng lẻ các công việc cá nhân.', style='List Bullet')
    
    # PHẦN II
    doc.add_heading('PHẦN II: QUY CHẾ ĐÁNH GIÁ VÀ CHẤM ĐIỂM CHI TIẾT', level=1)
    
    doc.add_heading('1. Nguyên tắc Không tính Tỷ trọng (Weighting) & Mắt xích đồng đẳng', level=2)
    doc.add_paragraph('Không áp dụng tỷ trọng %: Đặc thù các Ban dự án của chúng ta có chu kỳ công việc kéo dài, kết quả không rơi đều vào từng tháng. KPI ở đây đóng vai trò là "Công cụ kiểm soát tiến độ và tính kỷ luật", không phải là công cụ chia thưởng theo doanh thu. Việc gán tỷ trọng % sẽ rất khiên cưỡng và gây bất công nếu một việc bị tắc nghẽn bởi lý do khách quan.', style='List Bullet')
    p_dt = doc.add_paragraph(style='List Bullet')
    p_dt.add_run('Không phân loại công việc Định kỳ hay Phát sinh: ').bold = True
    p_dt.add_run('Hệ thống sẽ không phân loại trọng số điểm riêng biệt cho việc định kỳ (thường nhật) và việc phát sinh (do Sếp giao). Đã là công việc được phân rã từ Kế hoạch hoặc có Lệnh giao ban xuống thì mọi việc đều quan trọng như nhau.')
    doc.add_paragraph('Mắt xích đồng đẳng: Trong dự án, một công việc giấy tờ nhỏ đôi khi lại là điều kiện tiên quyết (nút thắt) để giải quyết một việc lớn. Do đó, mọi công việc đều được coi là những "mắt xích thiết yếu" bình đẳng. Đánh giá chỉ tập trung vào sự tuân thủ: Làm xong hay Trễ hạn.', style='List Bullet')
    
    doc.add_heading('2. Công thức tính điểm KPI (Cơ chế Điểm trừ)', level=2)
    doc.add_paragraph('Mặc định mỗi nhân sự đều có 100 điểm. Hoàn thành việc đúng hạn là trách nhiệm hiển nhiên để giữ điểm.')
    p = doc.add_paragraph()
    p.add_run('TỔNG ĐIỂM = 100 điểm - ĐIỂM PHẠT + ĐIỂM THƯỞNG').bold = True
    
    p = doc.add_paragraph()
    p.add_run('a. Khung Điểm Phạt (Trừ điểm)').bold = True
    doc.add_paragraph('Vi phạm nội quy, thái độ, đi trễ/về sớm không lý do liên tục: Trừ 5 điểm.', style='List Bullet')
    doc.add_paragraph('Không cập nhật báo cáo KPI trên phần mềm đúng hạn: Trừ 3 điểm/lần.', style='List Bullet')
    doc.add_paragraph('Trễ hạn deadline công việc: Trừ 2 điểm/ngày trễ.', style='List Bullet')
    doc.add_paragraph('Chất lượng không đạt (Bị trả về làm lại): Trừ 3 điểm/lần trả về.', style='List Bullet')
    doc.add_paragraph('Sai sót nghiệp vụ / Không hoàn thành tốt liên tục: Trừ 10 điểm (Trừ tối đa 10 điểm/task nếu không hoàn thành).', style='List Bullet')
    doc.add_paragraph('Trách nhiệm quản lý: Trừ 15 điểm đối với Trưởng ban/Quản lý nếu Chỉ tiêu chung không hoàn thành do lỗi điều phối, giao việc sai người hoặc không bám sát tiến độ.', style='List Bullet')
    
    p = doc.add_paragraph()
    p.add_run('b. Khung Điểm Thưởng (Cộng điểm)').bold = True
    doc.add_paragraph('Cộng thêm 5 điểm đối với các cá nhân hoàn thành xuất sắc nhiệm vụ: Hoàn thành các công việc trọng điểm liên kết với chỉ tiêu đúng hạn hoặc trước hạn (do Quản lý trực tiếp đánh giá và đề xuất).', style='List Bullet')
    
    doc.add_heading('3. Phân loại Hạng A*/A/B/C/D và Mức lương', level=2)
    doc.add_paragraph('Từ tổng điểm, hệ thống sẽ quy ra xếp loại lương vô cùng minh bạch:')
    doc.add_paragraph('Hạng A* (Vượt 100 điểm): Hưởng 110% đến 120% lương (Nhằm khích lệ sự vượt trội. Nếu duy trì A* nhiều quý, những nhân sự này sẽ được đề xuất nằm trong danh sách phê duyệt nhân viên xuất sắc của năm).', style='List Bullet')
    doc.add_paragraph('Hạng A: Hưởng 100% lương.', style='List Bullet')
    doc.add_paragraph('Hạng B & C: Hưởng từ 60% đến 80% lương.', style='List Bullet')
    doc.add_paragraph('Hạng D: Không đạt.', style='List Bullet')
    
    # PHẦN III
    doc.add_heading('PHẦN III: QUY TRÌNH DUYỆT VIỆC VÀ TÍNH NĂNG HỖ TRỢ', level=1)
    
    doc.add_paragraph('Vai trò của Quản lý trên phần mềm cực kỳ quyền lực thông qua 3 bước Duyệt trọng tâm:')
    
    doc.add_heading('1. Duyệt tiến độ hoàn thành Task (Kiểm soát chất lượng)', level=2)
    doc.add_paragraph('Khi nhân viên báo "Xong việc", hệ thống sẽ tự động thông báo cho Quản lý.', style='List Bullet')
    doc.add_paragraph('Quản lý bắt buộc phải truy cập vào xem file/minh chứng đính kèm. Chỉ khi minh chứng đạt chất lượng yêu cầu thì Quản lý mới bấm "Duyệt việc (Approve Task)". Nếu chưa đạt, có quyền trả về làm lại (Trừ 3 điểm KPI).', style='List Bullet')
    
    doc.add_heading('2. Duyệt việc khách quan (Bảo vệ quyền lợi nhân sự)', level=2)
    doc.add_paragraph('Nếu nhân viên gặp vướng mắc do yếu tố bên ngoài (chờ đối tác, thủ tục...), họ sẽ báo cáo lên phần mềm.', style='List Bullet')
    doc.add_paragraph('Quản lý truy cập vào mục "Duyệt việc khách quan" để đánh giá khối lượng thực tế nhân viên đã làm được và chốt tỷ lệ hoàn thành (VD: 50%, 80% hoặc miễn trừ).', style='List Bullet')
    doc.add_paragraph('Tỷ lệ % này liên kết trực tiếp vào điểm KPI cuối kỳ, đảm bảo quyền lợi công bằng nhất cho nhân viên.', style='List Bullet')
    
    doc.add_heading('3. Nghiệm thu Chỉ tiêu (Kiểm soát Master View)', level=2)
    doc.add_paragraph('Khi các nhân viên được giao việc thuộc một "Chỉ tiêu" báo cáo hoàn thành (đã được duyệt ở bước 1), thanh tiến độ (%) của Chỉ tiêu đó sẽ tự động chạy theo tỷ lệ tương ứng.', style='List Bullet')
    doc.add_paragraph('Tuy nhiên, Chỉ tiêu sẽ KHÔNG tự động chuyển sang trạng thái "Hoàn thành" dù tiến độ đã đạt 100%.', style='List Bullet')
    doc.add_paragraph('Quản lý phải là người đánh giá kết quả tổng thể cuối cùng và đích thân bấm nút "Nghiệm thu Chỉ tiêu (Close Objective)". Đây là bước chốt chặn để đảm bảo chất lượng công việc nhóm đã thực sự đáp ứng mục tiêu của Kế hoạch Năm/Quý đề ra.', style='List Bullet')
    
    doc.add_heading('Phân quyền và Bảo mật', level=2)
    doc.add_paragraph('Nhân viên: Giao diện tinh gọn, chỉ hiển thị đúng việc của cá nhân.', style='List Bullet')
    doc.add_paragraph('Quản lý: Bao quát toàn bộ tiến độ của phòng, sử dụng các nút quyền lực như Nghiệm thu Task, Duyệt việc khách quan, Nghiệm thu Chỉ tiêu.', style='List Bullet')
    doc.add_paragraph('HCNS / Ban Lãnh đạo: "Trạm kiểm soát trung tâm", nhìn được bức tranh tổng thể hoạt động của toàn công ty.', style='List Bullet')
    
    doc.add_heading('AI Rà soát Mô tả công việc (JD)', level=2)
    doc.add_paragraph('Hệ thống tích hợp Trí tuệ nhân tạo (AI) tự động đọc các công việc nhân viên đang làm thực tế và so sánh với bản Mô tả công việc (JD) ban đầu.', style='List Bullet')
    doc.add_paragraph('AI sẽ tự động cảnh báo nếu phát hiện nhân sự đang làm sai lệch chuyên môn, giúp Quản lý và HCNS kịp thời rà soát lại định biên nhân sự.', style='List Bullet')
    
    p = doc.add_paragraph()
    p.add_run('Lộ trình triển khai: Chạy thử nghiệm trong Quý IV/2026 để các Trưởng ban tập làm quen với kỹ năng "Rã việc" từ Kế hoạch Năm xuống.').italic = True
    
    doc.save('LoTrinh_QuyChe_KPI.docx')
    print("Done")

if __name__ == "__main__":
    main()
