from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()
    
    # Define colors
    dmt_blue = RGBColor(0, 51, 153) # Example corporate blue
    dark_gray = RGBColor(64, 64, 64)
    
    # ---------------------------------------------------
    # Slide 1: Title Slide
    # ---------------------------------------------------
    title_slide_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(title_slide_layout)
    title = slide1.shapes.title
    subtitle = slide1.placeholders[1]
    
    title.text = "CHUYỂN ĐỔI SỐ TRONG QUẢN TRỊ NHÂN SỰ\n& ĐÁNH GIÁ KPI"
    title.text_frame.paragraphs[0].font.color.rgb = dmt_blue
    title.text_frame.paragraphs[0].font.bold = True
    
    subtitle.text = ("Dự án: Hệ thống Quản trị & Đánh giá KPI Tự động\n"
                     "Đơn vị trình bày: Ban Hành chính Nhân sự\n"
                     "Mục tiêu: Minh bạch hóa năng suất - Tối ưu hóa nguồn lực")
    
    # ---------------------------------------------------
    # Slide 2: Problem Statement
    # ---------------------------------------------------
    bullet_slide_layout = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide2.shapes
    
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = "THỰC TRẠNG & NHỮNG BẤT CẬP HIỆN TẠI"
    title_shape.text_frame.paragraphs[0].font.color.rgb = dmt_blue
    
    tf = body_shape.text_frame
    tf.text = "Quy trình đánh giá hiện tại đang gặp các khó khăn:"
    
    p = tf.add_paragraph()
    p.text = "❌ Chấm điểm cảm tính: Phụ thuộc nhiều vào trí nhớ và đánh giá chủ quan."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "❌ Giao việc chồng chéo: Khó kiểm soát việc nhân sự làm sai chuyên môn (ngoài JD)."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "❌ Thiếu dữ liệu tổng quan: Ban Lãnh đạo không có cái nhìn tức thời (Real-time) về tỷ lệ hoàn thành công việc của từng phòng ban."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "❌ Mất nhiều thời gian: HCNS và Kế toán mất quá nhiều ngày để gom nhặt báo cáo, tính toán Excel dễ xảy ra sai sót."
    p.level = 1

    # ---------------------------------------------------
    # Slide 3: Solution
    # ---------------------------------------------------
    slide3 = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide3.shapes
    
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = "GIẢI PHÁP - HỆ THỐNG QUẢN TRỊ KPI MỚI"
    title_shape.text_frame.paragraphs[0].font.color.rgb = dmt_blue
    
    tf = body_shape.text_frame
    tf.text = "Thay đổi cách vận hành - Số hóa 100%"
    
    p = tf.add_paragraph()
    p.text = "✅ Chấm điểm tự động: Điểm số tính toán tự động dựa trên Khối lượng, Tiến độ (Gantt Chart) và Trọng số."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "✅ Tích hợp Trí tuệ Nhân tạo (AI): Tự động đối chiếu công việc thực tế với Mô tả công việc (JD), phát hiện ngay việc sai chuyên môn."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "✅ Dashboard Trực quan: Lãnh đạo theo dõi biểu đồ năng suất của từng Công ty/Phòng ban mọi lúc, mọi nơi."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "✅ Minh bạch & Công bằng: Cơ chế ghi nhận Điểm thưởng/Phạt rõ ràng trên hệ thống."
    p.level = 1

    # ---------------------------------------------------
    # Slide 4: Live Demo
    # ---------------------------------------------------
    slide4 = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide4.shapes
    
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = "TRẢI NGHIỆM THỰC TẾ (LIVE DEMO)"
    title_shape.text_frame.paragraphs[0].font.color.rgb = dmt_blue
    
    tf = body_shape.text_frame
    tf.text = "Các tính năng nổi bật:"
    
    p = tf.add_paragraph()
    p.text = "1️⃣ Biểu đồ Tổng quan (Dashboard): Số liệu tổng hợp ngay lập tức."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "2️⃣ Theo dõi Tiến độ (Gantt Chart): Trực quan hóa tiến độ dự án."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "3️⃣ Quét AI Hàng loạt (Magic Feature): Phát hiện nhân sự làm việc lặt vặt ngoài JD."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "4️⃣ Bảng tính Lương/Thưởng (KPI): Điểm số xuất tự động, không cần tính tay."
    p.level = 1

    # ---------------------------------------------------
    # Slide 5: Rollout Plan
    # ---------------------------------------------------
    slide5 = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide5.shapes
    
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    
    title_shape.text = "LỘ TRÌNH TRIỂN KHAI"
    title_shape.text_frame.paragraphs[0].font.color.rgb = dmt_blue
    
    tf = body_shape.text_frame
    tf.text = "Các bước đưa vào vận hành thực tế:"
    
    p = tf.add_paragraph()
    p.text = "Tháng 1 (Giai đoạn Chạy song song):"
    p.font.bold = True
    
    p = tf.add_paragraph()
    p.text = "Áp dụng thử nghiệm tại 2 phòng ban: HCNS & Kế toán."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Các phòng ban khác vẫn áp dụng quy chế cũ."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Tháng 2 (Giai đoạn Áp dụng toàn diện):"
    p.font.bold = True
    
    p = tf.add_paragraph()
    p.text = "Đóng hoàn toàn quy trình báo cáo Excel cũ."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Lấy 100% dữ liệu từ hệ thống mới làm cơ sở tính lương thưởng."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Cam kết: Tăng 40% hiệu suất quản lý báo cáo, giảm 100% sai sót do tính toán thủ công."
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 0, 0) # Highlight in red

    # Save presentation
    prs.save("Kế_Hoạch_Số_Hóa_KPI.pptx")
    print("Presentation saved as 'Kế_Hoạch_Số_Hóa_KPI.pptx'")

if __name__ == "__main__":
    create_presentation()
